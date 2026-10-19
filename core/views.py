from datetime import date, datetime, timedelta
from io import BytesIO
from pathlib import Path
import random
import re
from urllib.parse import urlparse

from django.conf import settings
from django.contrib import messages
from django.contrib.auth import authenticate, login, logout, update_session_auth_hash
from django.contrib.auth.decorators import login_required
from django.contrib.auth.forms import PasswordChangeForm
from django.contrib.auth.models import User
from django.contrib.auth.password_validation import validate_password
from django.core.exceptions import ValidationError
from django.core.validators import validate_email
from django.db import connection, transaction
from django.db.models import Count, Prefetch, Q
from django.http import FileResponse, Http404, HttpResponse, HttpResponseForbidden
from django.shortcuts import get_object_or_404, redirect, render
from django.utils import timezone
from django.utils._os import safe_join
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from .models import (Assignment, AuditLog, Course, Department, EmployeeType,
                     InternalDocument, Lesson, PasswordResetRequest, Profile, QualityDirection,
                     QualityRequirement, Question, Test, TestAttempt)


MAX_UPLOAD_SIZE = 15 * 1024 * 1024
ACTIVE_ASSIGNMENT_STATUSES = ('ASSIGNED', 'IN_PROGRESS', 'OVERDUE')


def health(request):
    """Small probe for the reverse proxy and container health check."""
    try:
        with connection.cursor() as cursor:
            cursor.execute('SELECT 1')
            cursor.fetchone()
    except Exception:
        return HttpResponse('unavailable', status=503, content_type='text/plain')
    return HttpResponse('ok', content_type='text/plain')


def profile(user): return Profile.objects.get_or_create(user=user)[0]
def role_in(user, *roles): return profile(user).role in roles
def can_manage(user): return role_in(user, 'ADMIN', 'METHODIST')
def can_manage_people(user): return role_in(user, 'ADMIN')
def is_head(user): return role_in(user, 'HEAD')
def deny(user): return HttpResponseForbidden('Недостаточно прав для этого раздела.')


def validate_pdf(upload):
    if not upload: return None
    if upload.size > MAX_UPLOAD_SIZE: return 'Размер PDF не должен превышать 15 МБ.'
    if Path(upload.name).suffix.lower() != '.pdf': return 'Разрешены только PDF-файлы.'
    if getattr(upload, 'content_type', '') not in ('application/pdf', 'application/octet-stream'):
        return 'Файл должен иметь тип PDF.'
    if upload.read(5) != b'%PDF-': return 'Содержимое файла не является PDF.'
    upload.seek(0)
    return None


def validate_drive_url(value):
    if not value:
        return 'Укажите ссылку Google Drive.'
    try:
        parsed = urlparse(value)
    except ValueError:
        return 'Некорректная ссылка Google Drive.'
    hostname = (parsed.hostname or '').lower()
    if parsed.scheme != 'https' or hostname not in ('drive.google.com', 'docs.google.com'):
        return 'Разрешены только HTTPS-ссылки drive.google.com или docs.google.com.'
    return None


@login_required
def media_file(request, path):
    try:
        filename = safe_join(settings.MEDIA_ROOT, path)
        return FileResponse(
            open(filename, 'rb'),
            as_attachment=request.GET.get('download') == '1',
            filename=Path(filename).name,
        )
    except (FileNotFoundError, ValueError):
        raise Http404


def login_view(request):
    if request.user.is_authenticated: return redirect('dashboard')
    error = None
    if request.method == 'POST':
        user = authenticate(request, username=request.POST.get('username'), password=request.POST.get('password'))
        if user and user.is_active:
            login(request, user)
            request.session.set_expiry(settings.SESSION_COOKIE_AGE)
            display_name = user.get_full_name().strip() or user.username
            messages.success(request, f'Добро пожаловать, {display_name}! Ваш логин: {user.username}.')
            AuditLog.objects.create(user=user, action='Вход в систему')
            return redirect('password' if profile(user).force_password_change else 'dashboard')
        error = 'Неверный логин или пароль.'
    return render(request, 'core/login.html', {'error': error})


def forgot_password(request):
    submitted = False
    error = None
    if request.method == 'POST':
        username = request.POST.get('username', '').strip()
        contact_method = request.POST.get('contact_method', '').strip()
        contact_value = request.POST.get('contact_value', '').strip()
        if contact_method not in dict(PasswordResetRequest.CONTACT_METHODS):
            error = 'Выберите WhatsApp или электронную почту.'
        elif contact_method == 'WHATSAPP':
            digits = re.sub(r'\D', '', contact_value)
            if not re.fullmatch(r'[\d\s()+-]+', contact_value) or not 10 <= len(digits) <= 15:
                error = 'Укажите номер WhatsApp полностью, например +7 777 123 45 67.'
            else:
                contact_value = '+' + digits
        else:
            try:
                validate_email(contact_value)
            except ValidationError:
                error = 'Укажите корректный адрес электронной почты.'
            else:
                contact_value = contact_value.lower()
        if error:
            return render(request, 'core/forgot_password.html', {
                'submitted': False, 'error': error, 'contact_method': contact_method,
                'contact_value': contact_value, 'username': username,
            })
        submitted = True
        user = User.objects.filter(
            username__iexact=username, is_active=True, profile__isnull=False,
        ).first()
        if user:
            password_request, created = PasswordResetRequest.objects.get_or_create(
                user=user,
                status='PENDING',
                defaults={
                    'contact_method': contact_method,
                    'contact_value': contact_value,
                },
            )
            # The latest contact is the one the employee can currently access.
            # Updating a pending request does not disclose whether the username exists.
            if not created and (
                password_request.contact_method != contact_method
                or password_request.contact_value != contact_value
            ):
                password_request.contact_method = contact_method
                password_request.contact_value = contact_value
                password_request.save(update_fields=['contact_method', 'contact_value'])
            if created:
                AuditLog.objects.create(
                    user=user, action='Запрошено восстановление пароля', object_label=user.username,
                )
    return render(request, 'core/forgot_password.html', {'submitted': submitted, 'error': error})


def logout_view(request): logout(request); return redirect('login')


@login_required
def session_ping(request):
    # Refresh older sessions that were created before the idle timeout changed.
    request.session.set_expiry(settings.SESSION_COOKIE_AGE)
    request.session.modified = True
    return HttpResponse(status=204)


@login_required
def change_password(request):
    form = PasswordChangeForm(request.user, request.POST or None)
    if request.method == 'POST' and form.is_valid():
        user = form.save(); update_session_auth_hash(request, user)
        p = profile(user); p.force_password_change = False; p.save(update_fields=['force_password_change'])
        return redirect('dashboard')
    return render(request, 'core/password.html', {'form': form})


def visible_assignments(user):
    qs = Assignment.objects.select_related(
        'user', 'user__profile', 'user__profile__department',
        'user__profile__employee_type', 'course', 'course__direction', 'test_version',
    ).prefetch_related('attempts')
    if role_in(user, 'ADMIN', 'METHODIST'): return qs
    if is_head(user): return qs.filter(user__profile__department=profile(user).department)
    return qs.filter(user=user)


def filtered_assignments(request, include_cancelled=False):
    assignments = visible_assignments(request.user)
    if not include_cancelled:
        assignments = assignments.exclude(status='CANCELLED')
    for field, lookup in [('direction','course__direction_id'),('requirement','course__quality_requirements__id'),('course','course_id'),('department','user__profile__department_id'),('employee_type','user__profile__employee_type_id'),('employee','user_id')]:
        value = request.GET.get(field)
        if value: assignments = assignments.filter(**{lookup: value})
    if request.GET.get('date_from'): assignments = assignments.filter(assigned_at__date__gte=request.GET['date_from'])
    if request.GET.get('date_to'): assignments = assignments.filter(assigned_at__date__lte=request.GET['date_to'])
    return assignments.distinct()


@login_required
def dashboard(request):
    p = profile(request.user)
    if p.role == 'EMPLOYEE': return redirect('directions')
    people = Profile.objects.filter(user__is_active=True)
    if p.role == 'HEAD': people = people.filter(department=p.department)

    # A reassignment creates a new cycle. Count only the latest live cycle for
    # each employee and course here; all earlier cycles remain in the report.
    current = {}
    for assignment in visible_assignments(request.user).filter(user__is_active=True):
        current.setdefault((assignment.user_id, assignment.course_id), assignment)
    assignments = [
        assignment for assignment in current.values()
        if assignment.status not in ('CANCELLED', 'REASSIGNED')
    ]
    statuses = {'ASSIGNED': 0, 'IN_PROGRESS': 0, 'OVERDUE': 0, 'FAILED': 0, 'COMPLETED': 0}
    course_rows = {}
    for assignment in assignments:
        status = 'OVERDUE' if assignment.is_overdue else assignment.status
        statuses[status] += 1
        row = course_rows.setdefault(assignment.course_id, {
            'course': assignment.course, 'total': 0,
            'assigned': 0, 'in_progress': 0, 'overdue': 0,
            'failed': 0, 'completed': 0, 'assignments': [],
        })
        row['total'] += 1
        row[status.lower()] += 1
        attempts = list(assignment.attempts.all())
        row['assignments'].append({
            'assignment': assignment,
            'status': status,
            'last_attempt': attempts[-1] if attempts else None,
        })

    if p.role in ('ADMIN', 'METHODIST'):
        for course in Course.objects.filter(active=True).select_related('direction'):
            course_rows.setdefault(course.id, {
                'course': course, 'total': 0,
                'assigned': 0, 'in_progress': 0, 'overdue': 0,
                'failed': 0, 'completed': 0, 'assignments': [],
            })
    courses = sorted(course_rows.values(), key=lambda row: row['course'].title.casefold())
    for row in courses:
        row['completion_rate'] = round(row['completed'] * 100 / row['total']) if row['total'] else 0
    selected_course = request.GET.get('course', '')
    selected_row = next(
        (row for row in courses if str(row['course'].id) == selected_course), None
    )
    if selected_row:
        selected_row['assignments'].sort(key=lambda item: (
            (item['assignment'].user.get_full_name() or item['assignment'].user.username).casefold(),
            item['assignment'].user_id,
        ))
        selected_row['departments'] = sorted({
            (item['assignment'].user.profile.department_id,
             item['assignment'].user.profile.department.name)
            for item in selected_row['assignments']
            if item['assignment'].user.profile.department_id
        }, key=lambda department: department[1].casefold())
        selected_row['types'] = sorted({
            (item['assignment'].user.profile.employee_type_id,
             item['assignment'].user.profile.employee_type.name)
            for item in selected_row['assignments']
            if item['assignment'].user.profile.employee_type_id
        }, key=lambda employee_type: employee_type[1].casefold())
        department = request.GET.get('department', '')
        employee_type = request.GET.get('employee_type', '')
        status = request.GET.get('status', '')
        search = request.GET.get('q', '').strip().casefold()
        selected_row['display_assignments'] = [
            item for item in selected_row['assignments']
            if (not department or str(item['assignment'].user.profile.department_id) == department)
            and (not employee_type or str(item['assignment'].user.profile.employee_type_id) == employee_type)
            and (not status or item['status'] == status)
            and (not search or search in (
                item['assignment'].user.get_full_name() + ' ' + item['assignment'].user.username
            ).casefold())
        ]
        selected_row['display_count'] = len(selected_row['display_assignments'])
    people_with_courses = len({assignment.user_id for assignment in assignments})
    return render(request, 'core/dashboard.html', {
        'profile': p, 'people_count': people.count(),
        'people_with_courses': people_with_courses,
        'people_without_courses': max(0, people.count() - people_with_courses),
        'assignment_count': len(assignments), 'course_count': len(courses),
        'assigned_course_count': sum(row['total'] > 0 for row in courses),
        'statuses': statuses, 'course_rows': courses, 'selected_row': selected_row,
    })


@login_required
def directions(request):
    items = QualityDirection.objects.filter(active=True).prefetch_related('courses', 'requirements')
    return render(request, 'core/directions.html', {'directions': items})


@login_required
def direction_detail(request, pk):
    direction = get_object_or_404(QualityDirection.objects.prefetch_related(
        Prefetch('courses', queryset=Course.objects.filter(active=True), to_attr='active_courses')
    ), pk=pk, active=True)
    return render(request, 'core/direction.html', {'direction': direction})


@login_required
def courses(request):
    # One card per course: only the newest assignment cycle defines whether
    # the course is currently assigned. Cancelled/reassigned cycles and old
    # records must not make a course reappear in the employee's list.
    latest_by_course = {}
    latest_completed_by_course = {}
    assignments = Assignment.objects.filter(
        user=request.user,
        course__active=True,
    ).select_related('course', 'course__direction', 'test_version').prefetch_related('attempts')
    for assignment in assignments:
        latest_by_course.setdefault(assignment.course_id, assignment)
        if assignment.status == 'COMPLETED' and assignment.test_version_id:
            latest_completed_by_course.setdefault(assignment.course_id, assignment)

    active_assignments = []
    for assignment in latest_by_course.values():
        if assignment.status in ('CANCELLED', 'REASSIGNED', 'COMPLETED') or not assignment.test_version_id:
            continue
        attempts = list(assignment.attempts.all())
        assignment.latest_attempt = attempts[-1] if attempts else None
        assignment.display_status = 'Просрочено' if assignment.is_overdue else assignment.get_status_display()
        active_assignments.append(assignment)
    completed_assignments = []
    for assignment in latest_completed_by_course.values():
        attempts = list(assignment.attempts.all())
        assignment.latest_attempt = attempts[-1] if attempts else None
        assignment.display_status = assignment.get_status_display()
        completed_assignments.append(assignment)
    active_assignments.sort(key=lambda item: item.course.title.casefold())
    completed_assignments.sort(key=lambda item: item.course.title.casefold())
    selected_tab = 'completed' if request.GET.get('tab') == 'completed' else 'active'
    my_assignments = completed_assignments if selected_tab == 'completed' else active_assignments
    return render(request, 'core/courses.html', {
        'my_assignments': my_assignments,
        'selected_tab': selected_tab,
        'active_count': len(active_assignments),
        'completed_count': len(completed_assignments),
    })


@login_required
def course_manage(request):
    if not can_manage(request.user): return deny(request.user)
    courses = Course.objects.select_related('direction').prefetch_related('lessons', 'tests', 'internal_documents')
    return render(request, 'core/course_manage.html', {'courses': courses})


@login_required
def course_form(request, pk=None):
    if not can_manage(request.user): return deny(request.user)
    course = get_object_or_404(Course, pk=pk) if pk else None
    if request.method == 'POST':
        course = course or Course()
        category_title = request.POST.get('category', '').strip() or 'Без категории'
        category = QualityDirection.objects.filter(title__iexact=category_title).first()
        if not category:
            category = QualityDirection.objects.create(title=category_title, active=True)
        course.direction = category
        course.title = request.POST['title'].strip()
        course.description = request.POST.get('description', '').strip()
        course.active = request.POST.get('active') == 'on'
        course.save()
        AuditLog.objects.create(
            user=request.user,
            action='Курс создан' if pk is None else 'Курс обновлён',
            object_label=course.title,
        )
        messages.success(request, 'Курс сохранён. Теперь можно добавить материалы и тест.')
        return redirect('course_content', course_id=course.id)
    return render(request, 'core/course_form.html', {
        'course': course,
        'directions': QualityDirection.objects.filter(active=True),
    })


@login_required
def course_content(request, course_id):
    if not can_manage(request.user): return deny(request.user)
    visible_lessons = Lesson.objects.exclude(category='INTERNAL')
    course = get_object_or_404(Course.objects.select_related('direction').prefetch_related(
        Prefetch('lessons', queryset=visible_lessons, to_attr='visible_lessons'),
        'tests', 'internal_documents',
    ), pk=course_id)
    documents = InternalDocument.objects.filter(status='ACTIVE').select_related('direction').order_by('title')
    return render(request, 'core/course_content.html', {'course': course, 'documents': documents})


@login_required
def course_document_link(request, course_id):
    if not can_manage(request.user): return deny(request.user)
    course = get_object_or_404(Course, pk=course_id)
    if request.method == 'POST':
        documents = InternalDocument.objects.filter(
            pk__in=request.POST.getlist('documents'), status='ACTIVE'
        )
        course.internal_documents.add(*documents)
        if documents:
            AuditLog.objects.create(user=request.user, action='Документы подключены к курсу', object_label=course.title)
            messages.success(request, f'Документы добавлены в курс «{course.title}».')
    return redirect('course_content', course_id=course.id)


@login_required
def course_document_unlink(request, course_id, document_id):
    if not can_manage(request.user): return deny(request.user)
    course = get_object_or_404(Course, pk=course_id)
    document = get_object_or_404(InternalDocument, pk=document_id)
    if request.method == 'POST':
        course.internal_documents.remove(document)
        AuditLog.objects.create(user=request.user, action='Документ откреплён от курса', object_label=f'{course.title}: {document.title}')
    return redirect('course_content', course_id=course.id)


@login_required
def lesson_create(request, course_id):
    if not can_manage(request.user): return deny(request.user)
    course = get_object_or_404(Course, pk=course_id)
    if request.method == 'POST':
        uploads = request.FILES.getlist('files') or ([request.FILES['file']] if request.FILES.get('file') else [])
        for upload in uploads:
            error = validate_pdf(upload)
            if error: messages.error(request, f'{upload.name}: {error}'); continue
            Lesson.objects.create(course=course, title=Path(upload.name).stem, category=request.POST.get('category','LESSON'), kind='DOC', order=int(request.POST.get('order') or 1), file=upload)
        if not uploads:
            Lesson.objects.create(course=course, title=request.POST['title'].strip(), category=request.POST.get('category','LESSON'), kind=request.POST.get('kind','LINK'), order=int(request.POST.get('order') or 1), url=request.POST.get('url','').strip())
        AuditLog.objects.create(user=request.user, action='Добавлен материал курса', object_label=course.title)
    return redirect('course_content', course_id=course.id)


@login_required
def lesson_delete(request, pk):
    if not can_manage(request.user): return deny(request.user)
    lesson = get_object_or_404(Lesson, pk=pk); course = lesson.course
    if request.method == 'POST': lesson.delete()
    return redirect('course_content', course_id=course.id)


@login_required
def lesson_edit(request, pk):
    if not can_manage(request.user): return deny(request.user)
    lesson = get_object_or_404(Lesson, pk=pk)
    if request.method == 'POST':
        lesson.title=request.POST['title'].strip(); lesson.category=request.POST['category']; lesson.kind=request.POST['kind']; lesson.order=max(1,int(request.POST.get('order') or 1)); lesson.url=request.POST.get('url','').strip()
        upload=request.FILES.get('file')
        if upload:
            error=validate_pdf(upload)
            if error: messages.error(request,error); return redirect('lesson_edit',pk=pk)
            lesson.file=upload
        lesson.save(); return redirect('course_content', course_id=lesson.course_id)
    return render(request,'core/lesson_form.html',{'lesson':lesson})


def current_test(course, published_only=False):
    qs = course.tests.filter(is_current=True)
    if published_only: qs = qs.filter(status='PUBLISHED')
    return qs.first()


def editable_test(course):
    test = current_test(course)
    if not test:
        return Test.objects.create(course=course, title=f'Итоговый тест: {course.title}', version=1)
    if test.status == 'PUBLISHED' and TestAttempt.objects.filter(assignment__test_version=test).exists():
        with transaction.atomic():
            test.is_current = False; test.save(update_fields=['is_current'])
            clone = Test.objects.create(course=course, title=test.title, instructions=test.instructions, time_limit_minutes=test.time_limit_minutes, version=test.version + 1, status='DRAFT', is_current=True)
            for q in test.questions.all():
                Question.objects.create(test=clone, text=q.text, question_type=q.question_type, options=q.options, correct_index=q.correct_index, correct_indexes=q.correct_indexes, points=q.points, explanation=q.explanation, order=q.order)
            return clone
    return test


@login_required
def test_editor(request, course_id):
    if not can_manage(request.user): return deny(request.user)
    course = get_object_or_404(Course, pk=course_id)
    test = current_test(course) or Test.objects.create(course=course, title=f'Итоговый тест: {course.title}')
    if request.method == 'POST':
        action = request.POST.get('action')
        if action == 'publish':
            if not test.questions.exists(): messages.error(request, 'Добавьте хотя бы один вопрос.')
            else:
                test.status = 'PUBLISHED'; test.save(update_fields=['status']); messages.success(request, f'Опубликована версия {test.version}.')
        elif action in ('save_test', 'save_question'):
            source_question = Question.objects.filter(pk=request.POST.get('question_id')).first() if request.POST.get('question_id') else None
            test = editable_test(course)
            if action == 'save_test':
                try:
                    time_limit_minutes = max(0, min(240, int(request.POST.get('time_limit_minutes', 0))))
                except (TypeError, ValueError):
                    time_limit_minutes = 0
                test.title = request.POST['title'].strip(); test.instructions = request.POST.get('instructions','').strip(); test.time_limit_minutes = time_limit_minutes; test.save()
            else:
                question = test.questions.filter(order=source_question.order, text=source_question.text).first() if source_question else Question(test=test)
                if question is None: question = Question(test=test)
                qtype = request.POST.get('question_type','SINGLE')
                options = ['Верно','Неверно'] if qtype == 'TRUE_FALSE' else [x.strip() for x in request.POST.get('options','').splitlines() if x.strip()]
                correct = [int(x) for x in request.POST.getlist('correct_indexes') if x.isdigit()]
                if qtype != 'MULTIPLE' and correct: correct = correct[:1]
                if not options or not correct or any(x >= len(options) for x in correct):
                    messages.error(request, 'Заполните варианты и выберите корректный ответ.')
                else:
                    question.text=request.POST['text'].strip(); question.question_type=qtype; question.options=options; question.correct_indexes=correct; question.correct_index=correct[0]; question.points=max(1,int(request.POST.get('points') or 1)); question.explanation=request.POST.get('explanation','').strip(); question.order=int(request.POST.get('order') or test.questions.count()+1); question.save()
        return redirect('test_editor', course_id=course.id)
    return render(request, 'core/test_editor.html', {'course': course, 'test': test, 'types': Question.TYPES})


@login_required
def test_preview(request, course_id):
    if not can_manage(request.user): return deny(request.user)
    course=get_object_or_404(Course,pk=course_id); test=current_test(course)
    return render(request,'core/test_preview.html',{'course':course,'test':test})


@login_required
def question_delete(request, pk):
    if not can_manage(request.user): return deny(request.user)
    source = get_object_or_404(Question, pk=pk); course = source.test.course
    if request.method == 'POST':
        test = editable_test(course)
        target = test.questions.filter(order=source.order, text=source.text).first()
        if target: target.delete()
    return redirect('test_editor', course_id=course.id)


@login_required
def course_detail(request, pk):
    visible_lessons = Lesson.objects.exclude(category='INTERNAL')
    course = get_object_or_404(
        Course.objects.prefetch_related(
            Prefetch('lessons', queryset=visible_lessons, to_attr='visible_lessons'),
            'internal_documents',
        ),
        pk=pk,
    )
    assignments = Assignment.objects.filter(course=course, user=request.user).prefetch_related('attempts')
    assignment = (
        assignments.filter(status__in=ACTIVE_ASSIGNMENT_STATUSES).first()
        or assignments.exclude(status__in=('CANCELLED', 'REASSIGNED')).first()
    )
    test = assignment.test_version if assignment and assignment.test_version_id else current_test(course, True)
    return render(request, 'core/course.html', {'course': course, 'assignment': assignment, 'test': test})


@login_required
def acknowledge(request, pk):
    assignment = get_object_or_404(Assignment, pk=pk, user=request.user)
    if request.method == 'POST':
        if assignment.status not in ACTIVE_ASSIGNMENT_STATUSES:
            messages.error(request, 'Назначение больше не активно.')
            return redirect('course', pk=assignment.course_id)
        if assignment.is_overdue:
            if assignment.status != 'OVERDUE':
                assignment.status = 'OVERDUE'; assignment.save(update_fields=['status'])
            messages.error(request, 'Срок прохождения курса истёк. Обратитесь к методисту для нового назначения.')
            return redirect('course', pk=assignment.course_id)
        assignment.acknowledged_at = timezone.now(); assignment.status = 'IN_PROGRESS'; assignment.save()
        AuditLog.objects.create(user=request.user, action='Начато обучение', object_label=assignment.course.title)
    return redirect('course', pk=assignment.course_id)


def test_assignment(course, user):
    assignments = Assignment.objects.filter(course=course, user=user)
    return (
        assignments.filter(status__in=ACTIVE_ASSIGNMENT_STATUSES).first()
        or assignments.exclude(status__in=('CANCELLED', 'REASSIGNED')).first()
    )


def test_timer_key(assignment):
    return f'test_started_at_{assignment.id}'


def test_started_at(request, assignment):
    # The database, not the browser session, owns the attempt deadline.
    return assignment.current_attempt_started_at


def test_seconds_left(test, started_at):
    if not test.time_limit_minutes or not started_at:
        return None
    return max(0, int(test.time_limit_minutes * 60 - (timezone.now() - started_at).total_seconds()))


def record_timeout(request, assignment, test):
    started_at = assignment.current_attempt_started_at
    if not started_at or test_seconds_left(test, started_at) != 0:
        return None
    with transaction.atomic():
        claimed = Assignment.objects.filter(
            pk=assignment.pk, current_attempt_started_at=started_at,
        ).update(current_attempt_started_at=None)
        if not claimed:
            return None
        assignment.refresh_from_db()
        attempt = TestAttempt.objects.create(
            assignment=assignment, number=assignment.attempts.count() + 1,
            correct_answers=0, score=0, passed=False,
            answers_snapshot=[{'timed_out': True}],
        )
        assignment.status = 'FAILED' if assignment.attempts_left == 0 else 'IN_PROGRESS'
        assignment.save(update_fields=['status'])
    request.session.pop(test_timer_key(assignment), None)
    AuditLog.objects.create(user=request.user, action='Время теста истекло', object_label=test.title)
    return attempt


@login_required
def start_test(request, pk):
    if request.method != 'POST':
        return redirect('test', pk=pk)
    course = get_object_or_404(Course, pk=pk)
    assignment = test_assignment(course, request.user)
    if not assignment or assignment.is_overdue or assignment.status == 'COMPLETED' or assignment.attempts_left == 0:
        return redirect('course', pk=pk)
    test = assignment.test_version or current_test(course, True)
    if not test:
        return HttpResponse('Итоговый тест ещё не опубликован.', status=409)
    started_at = assignment.current_attempt_started_at
    if started_at and test_seconds_left(test, started_at) == 0:
        record_timeout(request, assignment, test)
        return redirect('test', pk=pk)
    if not started_at:
        Assignment.objects.filter(
            pk=assignment.pk, current_attempt_started_at__isnull=True,
        ).update(current_attempt_started_at=timezone.now())
    request.session.pop(test_timer_key(assignment), None)
    return redirect('test', pk=pk)


@login_required
def take_test(request, pk):
    course = get_object_or_404(Course, pk=pk)
    assignment = test_assignment(course, request.user)
    if not assignment: return deny(request.user)
    if assignment.is_overdue:
        if assignment.status != 'OVERDUE':
            assignment.status = 'OVERDUE'; assignment.save(update_fields=['status'])
        messages.error(request, 'Срок теста истёк. Прохождение заблокировано.')
        return redirect('course', pk=pk)
    test = assignment.test_version or current_test(course, True)
    if not test: return HttpResponse('Итоговый тест ещё не опубликован.', status=409)
    if assignment.status == 'COMPLETED' or assignment.attempts_left == 0: return redirect('course', pk=pk)
    started_at = test_started_at(request, assignment)
    if not started_at:
        return render(request, 'core/test_start.html', {'test': test, 'assignment': assignment})
    seconds_left = test_seconds_left(test, started_at)
    if seconds_left == 0:
        attempt = record_timeout(request, assignment, test)
        if attempt is None:
            return redirect('test', pk=pk)
        return render(request, 'core/test.html', {'test': test, 'assignment': assignment, 'result': attempt, 'timed_out': True})
    if request.method == 'POST':
        questions=list(test.questions.all()); total=sum(q.points for q in questions); earned=0; correct_count=0; snapshot=[]
        for q in questions:
            selected = request.POST.getlist(f'q{q.id}')
            selected_indexes = sorted(int(x) for x in selected if x.isdigit())
            expected = sorted(q.correct_indexes or [q.correct_index])
            correct = selected_indexes == expected
            if correct: earned += q.points; correct_count += 1
            snapshot.append({'question_id':q.id,'question':q.text,'type':q.question_type,'options':q.options,'selected':selected_indexes,'correct':expected,'is_correct':correct,'points':q.points if correct else 0,'max_points':q.points})
        score = round(100 * earned / total) if total else 0; passed = score >= assignment.passing_score
        if test_seconds_left(test, started_at) == 0:
            attempt = record_timeout(request, assignment, test)
            if attempt is None:
                return redirect('test', pk=pk)
            return render(request, 'core/test.html', {'test':test, 'assignment':assignment, 'result':attempt, 'timed_out':True})
        with transaction.atomic():
            claimed = Assignment.objects.filter(
                pk=assignment.pk, current_attempt_started_at=started_at,
                status__in=ACTIVE_ASSIGNMENT_STATUSES,
            ).update(current_attempt_started_at=None)
            if not claimed:
                return redirect('test', pk=pk)
            assignment.refresh_from_db()
            attempt=TestAttempt.objects.create(assignment=assignment, number=assignment.attempts.count()+1, correct_answers=correct_count, score=score, passed=passed, answers_snapshot=snapshot)
            if passed: assignment.status='COMPLETED'; assignment.completed_at=timezone.now()
            elif assignment.attempts_left == 0: assignment.status='FAILED'
            else: assignment.status='IN_PROGRESS'
            assignment.save(update_fields=['status', 'completed_at'])
            AuditLog.objects.create(user=request.user, action='Пройден тест', object_label=f'{course.title}: {score}%')
        request.session.pop(test_timer_key(assignment), None)
        return render(request, 'core/test.html', {'test':test,'assignment':assignment,'result':attempt})
    question_items = []
    for question in test.questions.all():
        options = list(enumerate(question.options))
        random.SystemRandom().shuffle(options)
        question_items.append({'question': question, 'options': options})
    return render(request, 'core/test.html', {
        'test': test, 'assignment': assignment, 'question_items': question_items, 'seconds_left': seconds_left,
    })


@login_required
def document_bank(request):
    base_documents = InternalDocument.objects.select_related(
        'direction', 'requirement', 'department',
    ).prefetch_related('courses')
    all_documents = list(base_documents)
    documents = base_documents
    if request.GET.get('q'):
        query = request.GET['q'].strip()
        documents = documents.filter(
            Q(title__icontains=query) | Q(number__icontains=query)
            | Q(owner__icontains=query) | Q(tags__icontains=query)
        )
    for key, lookup in [('direction','direction_id'),('requirement','requirement_id'),('department','department_id'),('document_type','document_type'),('status','status')]:
        if request.GET.get(key): documents=documents.filter(**{lookup:request.GET[key]})
    if request.GET.get('owner'):
        documents = documents.filter(owner__icontains=request.GET['owner'].strip())
    review_filter = request.GET.get('review', '')
    today = timezone.localdate()
    review_limit = today + timedelta(days=30)
    if review_filter == 'OVERDUE':
        documents = documents.filter(status='ACTIVE', review_due__lt=today)
    elif review_filter == 'DUE_SOON':
        documents = documents.filter(status='ACTIVE', review_due__range=(today, review_limit))
    elif review_filter == 'NO_DATE':
        documents = documents.filter(status='ACTIVE', review_due__isnull=True)
    elif review_filter == 'CURRENT':
        documents = documents.filter(status='ACTIVE', review_due__gt=review_limit)

    active_documents = [doc for doc in all_documents if doc.status == 'ACTIVE']
    analytics = {
        'total': len(all_documents),
        'active': len(active_documents),
        'archived': sum(doc.status == 'ARCHIVED' for doc in all_documents),
        'overdue': sum(doc.review_state == 'OVERDUE' for doc in active_documents),
        'due_soon': sum(doc.review_state == 'DUE_SOON' for doc in active_documents),
        'no_review_date': sum(doc.review_state == 'NO_DATE' for doc in active_documents),
        'missing_owner': sum(not doc.owner.strip() for doc in active_documents),
        'owners': len({doc.owner.strip().casefold() for doc in active_documents if doc.owner.strip()}),
        'drive': sum(bool(doc.external_url) for doc in all_documents),
        'file': sum(bool(doc.file) for doc in all_documents),
    }
    owners = {}
    for doc in all_documents:
        owner = doc.owner.strip() or 'Не указан'
        owner_key = owner.casefold()
        row = owners.setdefault(owner_key, {
            'owner': owner, 'total': 0, 'active': 0, 'archived': 0,
            'overdue': 0, 'due_soon': 0,
        })
        row['total'] += 1
        row['active'] += doc.status == 'ACTIVE'
        row['archived'] += doc.status == 'ARCHIVED'
        row['overdue'] += doc.review_state == 'OVERDUE'
        row['due_soon'] += doc.review_state == 'DUE_SOON'
    owner_rows = sorted(
        owners.values(), key=lambda row: (-row['overdue'], -row['due_soon'], row['owner'].casefold()),
    )
    review_documents = sorted(
        [doc for doc in active_documents if doc.review_state in ('OVERDUE', 'DUE_SOON', 'NO_DATE')],
        key=lambda doc: (doc.review_due is None, doc.review_due or date.max, doc.title.casefold()),
    )
    document_view = request.GET.get('view', 'documents')
    if document_view not in ('documents', 'analytics'):
        document_view = 'documents'
    return render(request, 'core/document_bank.html', {
        'documents': documents.order_by('review_due', 'title'),
        'analytics': analytics,
        'owner_rows': owner_rows,
        'review_documents': review_documents,
        'document_view': document_view,
        'directions': QualityDirection.objects.all(),
        'requirements': QualityRequirement.objects.all(),
        'departments': Department.objects.all(),
        'document_types': InternalDocument.TYPES,
        'can_manage': can_manage(request.user),
        'selected_direction': request.GET.get('direction', ''),
        'selected_department': request.GET.get('department', ''),
        'selected_type': request.GET.get('document_type', ''),
        'selected_status': request.GET.get('status', ''),
        'selected_review': review_filter,
    })


@login_required
def document_form(request, pk=None):
    if not can_manage(request.user): return deny(request.user)
    document=get_object_or_404(InternalDocument,pk=pk) if pk else None
    if request.method=='POST':
        upload=request.FILES.get('file')
        source_type=request.POST.get('source_type','UPLOAD')
        external_url=request.POST.get('external_url','').strip()
        if upload:
            error=validate_pdf(upload)
            if error: messages.error(request,error); return redirect('document_edit',pk=pk) if pk else redirect('document_new')
        if source_type == 'DRIVE':
            error = validate_drive_url(external_url)
            if error:
                messages.error(request, error)
                return redirect('document_edit', pk=pk) if pk else redirect('document_new')
        elif source_type == 'UPLOAD':
            if not upload and not (document and document.file):
                messages.error(request, 'Выберите PDF-файл или источник Google Drive.')
                return redirect('document_edit', pk=pk) if pk else redirect('document_new')
        else:
            return HttpResponse('Некорректный источник документа.', status=400)
        was_new = document is None
        document=document or InternalDocument()
        for field in ('title','number','version','owner','tags','status','document_type'): setattr(document,field,request.POST.get(field,'').strip())
        document.direction_id=request.POST.get('direction') or None
        if 'requirement' in request.POST: document.requirement_id=request.POST.get('requirement') or None
        document.department_id=request.POST.get('department') or None
        document.approved_at=request.POST.get('approved_at') or None; document.review_due=request.POST.get('review_due') or None
        if source_type == 'DRIVE':
            document.external_url = external_url
            document.file = ''
        else:
            document.external_url = ''
            if upload: document.file=upload
        document.save()
        document.courses.set(request.POST.getlist('courses'))
        AuditLog.objects.create(
            user=request.user,
            action='Внутренний документ создан' if was_new else 'Внутренний документ обновлён',
            object_label=f'{document.number} {document.title}'.strip(),
        )
        messages.success(request, 'Документ сохранён в банке.')
        return redirect('document_bank')
    return render(request,'core/document_form.html',{'document':document,'directions':QualityDirection.objects.all(),'requirements':QualityRequirement.objects.all(),'departments':Department.objects.all(),'courses':Course.objects.all(),'document_types':InternalDocument.TYPES})


@login_required
def employees(request):
    if not can_manage_people(request.user): return deny(request.user)
    return render(request, 'core/employees.html', {'employees': Profile.objects.select_related('user','department','employee_type').order_by('user__last_name'), 'departments': Department.objects.all(), 'types': EmployeeType.objects.all()})


@login_required
def departments(request):
    if not can_manage_people(request.user): return deny(request.user)
    items = Department.objects.annotate(
        staff_count=Count('profile', distinct=True),
        document_count=Count('internal_documents', distinct=True),
    ).order_by('name')
    return render(request, 'core/departments.html', {'departments': items})


@login_required
def department_create(request):
    if not can_manage_people(request.user): return deny(request.user)
    if request.method != 'POST': return redirect('departments')
    name = request.POST.get('name', '').strip()
    if not name:
        messages.error(request, 'Введите название подразделения.')
    elif len(name) > Department._meta.get_field('name').max_length:
        messages.error(request, 'Название подразделения слишком длинное.')
    elif Department.objects.filter(name__iexact=name).exists():
        messages.error(request, 'Подразделение с таким названием уже существует.')
    else:
        department = Department.objects.create(name=name)
        AuditLog.objects.create(
            user=request.user, action='Создано подразделение', object_label=department.name,
        )
        messages.success(request, 'Подразделение создано.')
    return redirect('departments')


@login_required
def department_edit(request, pk):
    if not can_manage_people(request.user): return deny(request.user)
    department = get_object_or_404(Department, pk=pk)
    if request.method != 'POST': return redirect('departments')
    name = request.POST.get('name', '').strip()
    if not name:
        messages.error(request, 'Введите название подразделения.')
    elif len(name) > Department._meta.get_field('name').max_length:
        messages.error(request, 'Название подразделения слишком длинное.')
    elif Department.objects.filter(name__iexact=name).exclude(pk=department.pk).exists():
        messages.error(request, 'Такое подразделение уже существует. Используйте объединение.')
    else:
        old_name = department.name
        department.name = name
        department.save(update_fields=['name'])
        AuditLog.objects.create(
            user=request.user, action='Переименовано подразделение',
            object_label=f'{old_name} → {name}',
        )
        messages.success(request, 'Название подразделения обновлено.')
    return redirect('departments')


@login_required
def department_merge(request, pk):
    if not can_manage_people(request.user): return deny(request.user)
    if request.method != 'POST': return redirect('departments')
    source = get_object_or_404(Department, pk=pk)
    target = get_object_or_404(Department, pk=request.POST.get('target'))
    if source.pk == target.pk:
        messages.error(request, 'Нельзя объединить подразделение с самим собой.')
        return redirect('departments')
    with transaction.atomic():
        source = Department.objects.select_for_update().get(pk=source.pk)
        target = Department.objects.select_for_update().get(pk=target.pk)
        staff_count = Profile.objects.filter(department=source).update(department=target)
        document_count = InternalDocument.objects.filter(department=source).update(department=target)
        if not target.head_id and source.head_id:
            target.head_id = source.head_id
            target.save(update_fields=['head'])
        source_name = source.name
        target_name = target.name
        source.delete()
        AuditLog.objects.create(
            user=request.user, action='Объединены подразделения',
            object_label=f'{source_name} → {target_name}; сотрудников: {staff_count}',
        )
    messages.success(
        request,
        f'Подразделения объединены. Перенесено сотрудников: {staff_count}, документов: {document_count}.',
    )
    return redirect('departments')


@login_required
def department_delete(request, pk):
    if not can_manage_people(request.user): return deny(request.user)
    if request.method != 'POST': return redirect('departments')
    with transaction.atomic():
        department = get_object_or_404(Department.objects.select_for_update(), pk=pk)
        staff_count = Profile.objects.filter(department=department).count()
        document_count = InternalDocument.objects.filter(department=department).count()
        if staff_count or document_count or department.head_id:
            messages.error(
                request,
                'Нельзя удалить непустое подразделение. Сначала перенесите сотрудников '
                'и документы через объединение и снимите руководителя.',
            )
            return redirect('departments')
        department_name = department.name
        department.delete()
        AuditLog.objects.create(
            user=request.user, action='Удалено подразделение', object_label=department_name,
        )
    messages.success(request, f'Подразделение «{department_name}» удалено.')
    return redirect('departments')


@login_required
def employee_form(request, pk=None):
    if not can_manage_people(request.user): return deny(request.user)
    target = get_object_or_404(Profile, pk=pk) if pk else None
    if request.method == 'POST':
        username=request.POST['username'].strip(); user=target.user if target else User(username=username)
        user.username=username; user.first_name=request.POST['first_name'].strip(); user.last_name=request.POST['last_name'].strip()
        if not target: user.set_password(request.POST.get('password') or User.objects.make_random_password()); user.save(); target=Profile(user=user,force_password_change=True)
        else: user.save()
        target.role=request.POST['role']; target.department_id=request.POST['department']; target.employee_type_id=request.POST['employee_type']; target.position=request.POST['position'].strip(); target.employee_number=request.POST.get('employee_number','').strip(); target.phone=request.POST.get('phone','').strip(); target.save()
        AuditLog.objects.create(user=request.user, action='Профиль сотрудника сохранён', object_label=user.username)
        messages.success(request, 'Профиль сотрудника сохранён.')
        return redirect('employees')
    return render(request,'core/employee_form.html',{'employee':target,'departments':Department.objects.all(),'types':EmployeeType.objects.all(),'roles':Profile.ROLES})


@login_required
def password_requests(request):
    if not can_manage_people(request.user): return deny(request.user)
    requests = PasswordResetRequest.objects.filter(status='PENDING').select_related(
        'user', 'user__profile', 'user__profile__department', 'user__profile__employee_type',
    )
    return render(request, 'core/password_requests.html', {'password_requests': requests})


@login_required
def reset_employee_password(request, pk):
    if not can_manage_people(request.user): return deny(request.user)
    target = get_object_or_404(Profile.objects.select_related('user'), pk=pk)
    if request.method != 'POST': return redirect('employee_edit', pk=pk)
    if target.user_id == request.user.id:
        messages.error(request, 'Для смены собственного пароля используйте страницу смены пароля.')
        return redirect('employee_edit', pk=pk)
    password1 = request.POST.get('new_password1', '')
    password2 = request.POST.get('new_password2', '')
    if password1 != password2:
        messages.error(request, 'Введённые пароли не совпадают.')
        return redirect('employee_edit', pk=pk)
    try:
        validate_password(password1, target.user)
    except ValidationError as error:
        for message in error.messages: messages.error(request, message)
        return redirect('employee_edit', pk=pk)
    target.user.set_password(password1)
    target.user.save(update_fields=['password'])
    target.force_password_change = True
    target.save(update_fields=['force_password_change'])
    PasswordResetRequest.objects.filter(user=target.user, status='PENDING').update(
        status='COMPLETED', resolved_at=timezone.now(), resolved_by=request.user,
    )
    AuditLog.objects.create(
        user=request.user, action='Выдан временный пароль', object_label=target.user.username,
    )
    messages.success(request, 'Временный пароль установлен. При следующем входе сотрудник должен его изменить.')
    return redirect('employee_edit', pk=pk)


@login_required
def dismiss_password_request(request, pk):
    if not can_manage_people(request.user): return deny(request.user)
    password_request = get_object_or_404(PasswordResetRequest, pk=pk, status='PENDING')
    if request.method == 'POST':
        password_request.status = 'DISMISSED'
        password_request.resolved_at = timezone.now()
        password_request.resolved_by = request.user
        password_request.save(update_fields=['status', 'resolved_at', 'resolved_by'])
        AuditLog.objects.create(
            user=request.user, action='Запрос пароля закрыт без смены',
            object_label=password_request.user.username,
        )
        messages.success(request, 'Запрос закрыт без смены пароля.')
    return redirect('password_requests')


@login_required
def assignment_list(request):
    if not can_manage(request.user): return deny(request.user)
    return render(request,'core/assignments.html',{
        'assignments': visible_assignments(request.user),
        'courses': Course.objects.filter(active=True).select_related('direction').order_by('direction__title', 'title'),
        'departments': Department.objects.order_by('name'),
        'types': EmployeeType.objects.order_by('name'),
        'employees': Profile.objects.select_related('user','department','employee_type').filter(user__is_active=True).order_by('user__last_name', 'user__first_name'),
        'selected_course': request.GET.get('course',''),
    })


@login_required
def create_assignment(request):
    if not can_manage(request.user): return deny(request.user)
    if request.method!='POST': return redirect('assignments')
    course=get_object_or_404(Course,pk=request.POST['course']); test=current_test(course,True)
    if not test: return HttpResponse('Сначала опубликуйте итоговый тест курса.',status=409)
    due=datetime.strptime(request.POST['due_date'],'%Y-%m-%d').date()
    due_time=datetime.strptime(request.POST.get('due_time') or '23:59','%H:%M').time()
    requested_user_ids = [value for value in request.POST.getlist('users') if value.isdigit()]
    # The browser list and the server use the same source of truth.  A modified
    # POST therefore cannot assign a course to an inactive account or a user
    # without an employee profile.
    user_ids = set(
        Profile.objects.filter(
            user__is_active=True, user_id__in=requested_user_ids,
        ).values_list('user_id', flat=True)
    )
    department_ids=request.POST.getlist('departments')
    employee_type_ids=request.POST.getlist('employee_types')
    audience = Profile.objects.filter(user__is_active=True)
    if department_ids: audience = audience.filter(department_id__in=department_ids)
    if employee_type_ids: audience = audience.filter(employee_type_id__in=employee_type_ids)
    if department_ids or employee_type_ids:
        user_ids.update(audience.values_list('user_id',flat=True))
    if not user_ids: return HttpResponse('Выберите хотя бы одну категорию, отделение или сотрудника.',status=400)
    values={'due_date':due,'due_time':due_time,'passing_score':int(request.POST['passing_score']),'attempts_allowed':int(request.POST['attempts_allowed']),'test_version':test}
    for user_id in user_ids:
        Assignment.objects.filter(
            user_id=user_id, course=course,
            status__in=ACTIVE_ASSIGNMENT_STATUSES,
        ).update(status='REASSIGNED')
        Assignment.objects.create(user_id=user_id,course=course,**values)
    AuditLog.objects.create(user=request.user,action='Назначено обучение',object_label=course.title); return redirect('assignments')


@login_required
def assignment_bulk_action(request):
    if not can_manage(request.user): return deny(request.user)
    if request.method != 'POST': return redirect('assignments')
    course_ids = request.POST.getlist('courses')
    if not course_ids and request.POST.get('course'):
        course_ids = [request.POST['course']]
    try:
        courses = list(Course.objects.filter(pk__in=course_ids))
    except (TypeError, ValueError):
        courses = []
    if not courses or len(courses) != len(set(course_ids)):
        messages.error(request, 'Выберите хотя бы один курс для массовой операции.')
        return redirect('assignments')
    department_ids = request.POST.getlist('departments')
    employee_type_ids = request.POST.getlist('employee_types')
    all_audience = request.POST.get('all_audience') == 'on'
    if not all_audience and not department_ids and not employee_type_ids:
        messages.error(request, 'Выберите хотя бы одну категорию или отделение либо отметьте «Все сотрудники».')
        return redirect('assignments')
    eligible_statuses = list(ACTIVE_ASSIGNMENT_STATUSES)
    if request.POST.get('only_incomplete') == 'on':
        eligible_statuses.append('FAILED')
    assignment_scope = Assignment.objects.filter(
        course_id__in=course_ids, user__is_active=True,
    ).exclude(status__in=('CANCELLED', 'REASSIGNED'))
    if department_ids:
        assignment_scope = assignment_scope.filter(user__profile__department_id__in=department_ids)
    if employee_type_ids:
        assignment_scope = assignment_scope.filter(user__profile__employee_type_id__in=employee_type_ids)
    action = request.POST.get('action')
    with transaction.atomic():
        latest_ids = []
        seen_pairs = set()
        for assignment_id, user_id, course_id in assignment_scope.order_by(
            'user_id', 'course_id', '-assigned_at', '-id',
        ).values_list('id', 'user_id', 'course_id'):
            pair = (user_id, course_id)
            if pair not in seen_pairs:
                seen_pairs.add(pair)
                latest_ids.append(assignment_id)
        assignments = Assignment.objects.select_for_update().filter(
            pk__in=latest_ids, status__in=eligible_statuses,
        )
        selected = list(assignments.select_related('user', 'course').prefetch_related('attempts'))
        if not selected:
            messages.warning(request, 'Для выбранных курсов и аудитории незавершённых назначений не найдено.')
            return redirect('assignments')
        if action == 'cancel':
            now = timezone.now()
            for assignment in selected:
                assignment.status = 'CANCELLED'
                assignment.cancelled_at = now
                assignment.cancelled_by = request.user
            Assignment.objects.bulk_update(selected, ['status', 'cancelled_at', 'cancelled_by'])
            AuditLog.objects.create(
                user=request.user, action='Массовая отмена назначений',
                object_label=f'Курсов: {len(courses)}; назначений: {len(selected)}',
            )
            messages.success(request, f'Отменено назначений: {len(selected)}. Записи сохранены в журнале назначений и скрыты из отчёта.')
            return redirect('assignments')
        if action != 'edit':
            messages.error(request, 'Неизвестная массовая операция.')
            return redirect('assignments')
        try:
            due_date = datetime.strptime(request.POST['due_date'], '%Y-%m-%d').date()
            due_time = datetime.strptime(request.POST.get('due_time') or '23:59', '%H:%M').time()
            passing_score_raw = request.POST.get('passing_score', '').strip()
            attempts_allowed_raw = request.POST.get('attempts_allowed', '').strip()
            passing_score = int(passing_score_raw) if passing_score_raw else None
            attempts_allowed = int(attempts_allowed_raw) if attempts_allowed_raw else None
            additional_attempts = int(request.POST.get('additional_attempts') or 0)
        except (KeyError, TypeError, ValueError):
            messages.error(request, 'Проверьте дату, время, порог и количество попыток.')
            return redirect('assignments')
        if passing_score is not None and not 1 <= passing_score <= 100:
            messages.error(request, 'Порог должен быть от 1 до 100%.')
            return redirect('assignments')
        if attempts_allowed is not None and attempts_allowed < 1:
            messages.error(request, 'Количество попыток должно быть не меньше одной.')
            return redirect('assignments')
        if additional_attempts < 0:
            messages.error(request, 'Количество дополнительных попыток не может быть отрицательным.')
            return redirect('assignments')
        if attempts_allowed is not None and additional_attempts:
            messages.error(request, 'Укажите либо общее число попыток, либо количество дополнительных попыток.')
            return redirect('assignments')
        for assignment in selected:
            attempts_used = len(assignment.attempts.all())
            if attempts_used and passing_score is not None and passing_score != assignment.passing_score:
                messages.error(
                    request,
                    'Массовое изменение отменено: у части сотрудников уже есть попытки, '
                    'поэтому порог менять нельзя.',
                )
                return redirect('assignments')
            if attempts_allowed is not None and attempts_allowed < attempts_used:
                messages.error(
                    request,
                    'Массовое изменение отменено: новый лимит меньше числа уже использованных попыток.',
                )
                return redirect('assignments')
        now = timezone.now()
        for assignment in selected:
            attempts_used = len(assignment.attempts.all())
            assignment.due_date = due_date
            assignment.due_time = due_time
            if passing_score is not None:
                assignment.passing_score = passing_score
            if attempts_allowed is not None:
                assignment.attempts_allowed = attempts_allowed
            elif additional_attempts:
                assignment.attempts_allowed += additional_attempts
            if attempts_used and attempts_used >= assignment.attempts_allowed:
                assignment.status = 'FAILED'
            elif now > assignment.deadline:
                assignment.status = 'OVERDUE'
            else:
                assignment.status = 'IN_PROGRESS' if assignment.acknowledged_at else 'ASSIGNED'
        Assignment.objects.bulk_update(selected, [
            'due_date', 'due_time', 'passing_score', 'attempts_allowed', 'status',
        ])
        AuditLog.objects.create(
            user=request.user, action='Массовое изменение назначений',
            object_label=f'Курсов: {len(courses)}; назначений: {len(selected)}',
        )
        messages.success(request, f'Обновлено назначений: {len(selected)}.')
    return redirect('assignments')


@login_required
def assignment_edit(request, pk):
    if not can_manage(request.user): return deny(request.user)
    assignment = get_object_or_404(
        Assignment.objects.select_related('user', 'course').prefetch_related('attempts'), pk=pk,
    )
    if assignment.status not in ACTIVE_ASSIGNMENT_STATUSES:
        messages.error(request, 'Можно редактировать только активное назначение.')
        return redirect('assignments')
    attempts_used = assignment.attempts.count()
    if request.method == 'POST':
        try:
            due_date = datetime.strptime(request.POST['due_date'], '%Y-%m-%d').date()
            due_time = datetime.strptime(request.POST.get('due_time') or '23:59', '%H:%M').time()
            passing_score = int(request.POST['passing_score'])
            attempts_allowed = int(request.POST['attempts_allowed'])
        except (KeyError, TypeError, ValueError):
            messages.error(request, 'Проверьте дату, время, порог и количество попыток.')
            return redirect('assignment_edit', pk=assignment.id)
        if not 1 <= passing_score <= 100 or attempts_allowed < 1:
            messages.error(request, 'Порог должен быть от 1 до 100%, а количество попыток — не меньше одной.')
            return redirect('assignment_edit', pk=assignment.id)
        if attempts_used and passing_score != assignment.passing_score:
            messages.error(request, 'Порог нельзя изменить после начала тестирования: история оценок должна оставаться неизменной.')
            return redirect('assignment_edit', pk=assignment.id)
        if attempts_allowed < attempts_used:
            messages.error(request, f'Нельзя установить меньше {attempts_used} попыток: они уже использованы.')
            return redirect('assignment_edit', pk=assignment.id)
        assignment.due_date = due_date
        assignment.due_time = due_time
        assignment.passing_score = passing_score
        assignment.attempts_allowed = attempts_allowed
        if attempts_used and attempts_used >= attempts_allowed:
            assignment.status = 'FAILED'
        elif timezone.now() > assignment.deadline:
            assignment.status = 'OVERDUE'
        else:
            assignment.status = 'IN_PROGRESS' if assignment.acknowledged_at else 'ASSIGNED'
        assignment.save(update_fields=[
            'due_date', 'due_time', 'passing_score', 'attempts_allowed', 'status',
        ])
        AuditLog.objects.create(
            user=request.user, action='Назначение изменено',
            object_label=f'#{assignment.id} {assignment.user.username}: {assignment.course.title}',
        )
        messages.success(request, 'Назначение обновлено.')
        return redirect('assignments')
    return render(request, 'core/assignment_edit.html', {
        'assignment': assignment, 'attempts_used': attempts_used,
    })


@login_required
def assignment_cancel(request, pk):
    if not can_manage(request.user): return deny(request.user)
    assignment = get_object_or_404(Assignment.objects.select_related('user', 'course'), pk=pk)
    if request.method != 'POST': return redirect('assignments')
    if assignment.status not in ACTIVE_ASSIGNMENT_STATUSES:
        messages.error(request, 'Можно отменить только активное назначение.')
        return redirect('assignments')
    assignment.status = 'CANCELLED'
    assignment.cancelled_at = timezone.now()
    assignment.cancelled_by = request.user
    assignment.save(update_fields=['status', 'cancelled_at', 'cancelled_by'])
    AuditLog.objects.create(
        user=request.user, action='Назначение отменено',
        object_label=f'#{assignment.id} {assignment.user.username}: {assignment.course.title}',
    )
    messages.success(request, 'Назначение отменено. Запись сохранена в журнале назначений и скрыта из отчёта.')
    return redirect('assignments')


def report_entry(assignment):
    attempts = list(assignment.attempts.all())
    last_attempt = attempts[-1] if attempts else None
    status = 'OVERDUE' if assignment.is_overdue else assignment.status
    return {
        'assignment': assignment,
        'user': assignment.user,
        'profile': assignment.user.profile,
        'course': assignment.course,
        'last_attempt': last_attempt,
        'score': last_attempt.score if last_attempt else None,
        'result_at': last_attempt.completed_at if last_attempt else None,
        'attempts_used': len(attempts),
        'status': status,
        'status_label': dict(Assignment.STATUS).get(status, status),
    }


def report_sort(rows, sort, order):
    def name(row):
        user = row['user']
        return (user.get_full_name() or user.username).casefold()

    keys = {
        'employee': name,
        'course': lambda row: row['course'].title.casefold(),
        'assigned': lambda row: row['assignment'].assigned_at.timestamp(),
        'due': lambda row: row['assignment'].deadline.timestamp(),
        'status': lambda row: row['status_label'].casefold(),
        'score': lambda row: row['score'] if row['score'] is not None else -1,
        'completed': lambda row: row['result_at'].timestamp() if row['result_at'] else -1,
        'attempts': lambda row: row['attempts_used'],
    }
    return sorted(rows, key=keys.get(sort, name), reverse=order == 'desc')


def current_report_entries(assignments):
    latest = {}
    for assignment in sorted(assignments, key=lambda item: (item.assigned_at, item.id), reverse=True):
        latest.setdefault((assignment.user_id, assignment.course_id), assignment)
    return [
        report_entry(assignment) for assignment in latest.values()
        if assignment.status not in ('CANCELLED', 'REASSIGNED')
    ]


def status_filtered(entries, request):
    status = request.GET.get('status', '')
    return [entry for entry in entries if not status or entry['status'] == status]


def report_data(request):
    assignments = list(filtered_assignments(request, include_cancelled=True))
    history_entries = status_filtered([
        report_entry(assignment) for assignment in assignments
        if assignment.status != 'CANCELLED'
    ], request)
    current_entries = status_filtered(current_report_entries(assignments), request)

    sort = request.GET.get('sort', 'employee')
    order = request.GET.get('order', 'asc')
    if order not in ('asc', 'desc'):
        order = 'asc'
    history_rows = report_sort(history_entries, sort, order)

    selected_course_id = request.GET.get('course')
    selected_course = None
    if selected_course_id:
        selected_course = Course.objects.filter(pk=selected_course_id).first()
    if not selected_course and current_entries:
        selected_course = sorted(
            {entry['course'] for entry in current_entries},
            key=lambda course: course.title.casefold(),
        )[0]
    course_rows = [
        entry for entry in current_entries
        if selected_course and entry['course'].id == selected_course.id
    ]
    course_rows = report_sort(course_rows, sort, order)

    matrix_courses = sorted(
        {entry['course'] for entry in current_entries},
        key=lambda course: course.title.casefold(),
    )
    matrix_users = {}
    for entry in current_entries:
        matrix_users.setdefault(entry['user'].id, {
            'user': entry['user'], 'profile': entry['profile'], 'entries': {},
        })['entries'][entry['course'].id] = entry
    matrix_rows = []
    for row in sorted(
        matrix_users.values(),
        key=lambda item: (item['user'].get_full_name() or item['user'].username).casefold(),
    ):
        row['cells'] = [row['entries'].get(course.id) for course in matrix_courses]
        scores = [cell['score'] for cell in row['cells'] if cell and cell['score'] is not None]
        row['average'] = round(sum(scores) / len(scores), 1) if scores else None
        matrix_rows.append(row)

    analytics = []
    for course in matrix_courses:
        rows = [entry for entry in current_entries if entry['course'].id == course.id]
        scores = [entry['score'] for entry in rows if entry['score'] is not None]
        counts = {key: sum(entry['status'] == key for entry in rows) for key, _ in Assignment.STATUS}
        total = len(rows)
        analytics.append({
            'course': course,
            'total': total,
            'assigned': counts['ASSIGNED'],
            'in_progress': counts['IN_PROGRESS'],
            'completed': counts['COMPLETED'],
            'failed': counts['FAILED'],
            'overdue': counts['OVERDUE'],
            'average': round(sum(scores) / len(scores), 1) if scores else None,
            'completion_rate': round(counts['COMPLETED'] * 100 / total) if total else 0,
        })

    selected_scores = [row['score'] for row in course_rows if row['score'] is not None]
    selected_stats = {
        'total': len(course_rows),
        'assigned': sum(row['status'] == 'ASSIGNED' for row in course_rows),
        'in_progress': sum(row['status'] == 'IN_PROGRESS' for row in course_rows),
        'completed': sum(row['status'] == 'COMPLETED' for row in course_rows),
        'failed': sum(row['status'] == 'FAILED' for row in course_rows),
        'overdue': sum(row['status'] == 'OVERDUE' for row in course_rows),
        'average': round(sum(selected_scores) / len(selected_scores), 1) if selected_scores else None,
    }
    selected_stats['completion_rate'] = (
        round(selected_stats['completed'] * 100 / selected_stats['total'])
        if selected_stats['total'] else 0
    )
    return {
        'course_rows': course_rows,
        'history_rows': history_rows,
        'matrix_rows': matrix_rows,
        'matrix_courses': matrix_courses,
        'analytics': analytics,
        'selected_course': selected_course,
        'selected_stats': selected_stats,
        'sort': sort,
        'order': order,
    }


@login_required
def reports(request):
    if role_in(request.user,'EMPLOYEE'): return deny(request.user)
    data = report_data(request)
    report_view = request.GET.get('view', 'course')
    if report_view not in ('course', 'matrix', 'history'):
        report_view = 'course'
    base_query = request.GET.copy()
    for key in ('view', 'sort', 'order'):
        base_query.pop(key, None)
    employees = Profile.objects.select_related('user')
    if is_head(request.user):
        employees = employees.filter(department=profile(request.user).department)
    data.update({
        'report_view': report_view,
        'base_query': base_query.urlencode(),
        'directions': QualityDirection.objects.all(),
        'requirements': QualityRequirement.objects.all(),
        'courses': Course.objects.all(),
        'departments': Department.objects.all(),
        'employee_types': EmployeeType.objects.all(),
        'employees': employees,
        'status_choices': Assignment.STATUS,
        'selected_direction': request.GET.get('direction', ''),
        'selected_course_id': request.GET.get('course', ''),
        'selected_department': request.GET.get('department', ''),
        'selected_employee_type': request.GET.get('employee_type', ''),
        'selected_employee': request.GET.get('employee', ''),
        'selected_status': request.GET.get('status', ''),
        # Compatibility for callers that inspected the old context key.
        'report_rows': data['history_rows'],
    })
    return render(request, 'core/reports.html', data)


@login_required
def report_excel(request):
    if role_in(request.user,'EMPLOYEE'): return deny(request.user)
    data = report_data(request)
    wb = Workbook()
    journal = wb.active
    journal.title = 'Журнал назначений'
    journal.append([
        'Сотрудник', 'Логин', 'Категория сотрудника', 'Отдел / отделение', 'Должность',
        'Курс', 'Дата назначения', 'Срок прохождения', 'Статус', 'Последняя оценка, %',
        'Дата результата', 'Использовано попыток', 'Разрешено попыток', 'Порог, %',
    ])
    for entry in data['history_rows']:
        assignment = entry['assignment']
        user = entry['user']
        employee = entry['profile']
        journal.append([
            user.get_full_name() or user.username,
            user.username,
            str(employee.employee_type or ''),
            str(employee.department or ''),
            employee.position,
            entry['course'].title,
            timezone.localtime(assignment.assigned_at).replace(tzinfo=None),
            timezone.localtime(assignment.deadline).replace(tzinfo=None),
            entry['status_label'],
            entry['score'],
            timezone.localtime(entry['result_at']).replace(tzinfo=None) if entry['result_at'] else None,
            entry['attempts_used'],
            assignment.attempts_allowed,
            assignment.passing_score,
        ])

    matrix = wb.create_sheet('Матрица обучения')
    matrix.append([
        'Сотрудник', 'Логин', 'Категория сотрудника', 'Отдел / отделение', 'Должность',
        *[course.title for course in data['matrix_courses']], 'Средний балл, %',
    ])
    for row in data['matrix_rows']:
        user = row['user']
        employee = row['profile']
        matrix.append([
            user.get_full_name() or user.username,
            user.username,
            str(employee.employee_type or ''),
            str(employee.department or ''),
            employee.position,
            *[
                cell['score'] if cell and cell['score'] is not None
                else (cell['status_label'] if cell else None)
                for cell in row['cells']
            ],
            row['average'],
        ])

    analytics = wb.create_sheet('Аналитика по курсам')
    analytics.append([
        'Курс', 'Назначено сотрудникам', 'Ожидают начала', 'В процессе', 'Завершили',
        'Не прошли', 'Просрочили', 'Средний балл, %', 'Процент завершения, %',
    ])
    for row in data['analytics']:
        analytics.append([
            row['course'].title, row['total'], row['assigned'], row['in_progress'],
            row['completed'], row['failed'], row['overdue'], row['average'],
            row['completion_rate'],
        ])

    header_fill = PatternFill('solid', fgColor='103B53')
    for sheet in wb.worksheets:
        sheet.freeze_panes = 'A2'
        sheet.auto_filter.ref = sheet.dimensions
        sheet.sheet_view.showGridLines = False
        for cell in sheet[1]:
            cell.font = Font(color='FFFFFF', bold=True)
            cell.fill = header_fill
            cell.alignment = Alignment(wrap_text=True, vertical='center')
        sheet.row_dimensions[1].height = 34
        for column in sheet.columns:
            values = [len(str(cell.value or '')) for cell in column[:1000]]
            sheet.column_dimensions[column[0].column_letter].width = min(42, max(12, max(values, default=10) + 2))
    for row in range(2, journal.max_row + 1):
        journal.cell(row, 7).number_format = 'dd.mm.yyyy hh:mm'
        journal.cell(row, 8).number_format = 'dd.mm.yyyy hh:mm'
        journal.cell(row, 11).number_format = 'dd.mm.yyyy hh:mm'

    output = BytesIO()
    wb.save(output)
    response = HttpResponse(
        output.getvalue(),
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
    )
    response['Content-Disposition'] = 'attachment; filename="learning-results.xlsx"'
    return response
