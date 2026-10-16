from datetime import date, timedelta
from io import BytesIO
from django.conf import settings
from django.contrib.auth.models import User
from django.contrib.staticfiles import finders
from django.core.files.uploadedfile import SimpleUploadedFile
from django.core.management import call_command
from django.db.models.deletion import ProtectedError
from django.test import TestCase
from django.urls import reverse
from django.utils import timezone
from openpyxl import load_workbook
from .models import Assignment, Course, Department, EmployeeType, InternalDocument, Lesson, PasswordResetRequest, Profile, QualityDirection, Question, Test, TestAttempt
from .training_content import BLUE_CODE_ADULT_QUESTIONS

class LearningPortalTests(TestCase):
    def test_health_reports_database_availability(self):
        response = self.client.get(reverse('health'))
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.content, b'ok')
    def setUp(self):
        dep = Department.objects.create(name='Тестовое отделение'); kind = EmployeeType.objects.create(name='Медицинский')
        self.employee = User.objects.create_user('employee', password='test-password'); Profile.objects.create(user=self.employee, role='EMPLOYEE', department=dep, employee_type=kind, position='Врач')
        methodist = User.objects.create_user('methodist', password='test-password'); Profile.objects.create(user=methodist, role='METHODIST', department=dep, employee_type=kind, position='Методист')
        self.admin = User.objects.create_user('sysadmin', password='admin-password'); Profile.objects.create(user=self.admin, role='ADMIN', position='Системный администратор')
        direction = QualityDirection.objects.create(title='Безопасность пациента'); self.course = Course.objects.create(direction=direction, title='Тестовый курс')
        self.test = Test.objects.create(course=self.course, title='Итоговый тест', status='PUBLISHED'); Question.objects.create(test=self.test, text='Верный вариант?', options=['Да','Нет'], correct_index=0, correct_indexes=[0])
        self.assignment = Assignment.objects.create(user=self.employee, course=self.course, test_version=self.test, due_date=timezone.localdate()+timedelta(days=1), passing_score=80, attempts_allowed=1)
    def test_closed_dashboard(self): self.assertRedirects(self.client.get(reverse('dashboard')), '/quality/login/?next=/quality/')
    def test_login_and_portal_use_cyamio_branding(self):
        self.assertIn('whitenoise.middleware.WhiteNoiseMiddleware', settings.MIDDLEWARE)
        login_page = self.client.get(reverse('login'))
        self.assertContains(login_page, 'login-brand-logo')
        self.assertContains(login_page, 'core/cyamio-logo')
        self.assertContains(login_page, '.png')
        self.assertContains(login_page, 'Портал непрерывного повышения качества медицинских услуг')
        self.assertIsNotNone(finders.find('core/cyamio-logo.png'))

        self.client.login(username='sysadmin', password='admin-password')
        portal_page = self.client.get(reverse('dashboard'))
        self.assertContains(portal_page, 'sidebar-logo')
        self.assertContains(portal_page, 'Эмблема ЦЯМиО')
    def test_department_assignment_creates_personal_assignment(self):
        self.client.login(username='methodist', password='test-password'); response = self.client.post(reverse('assignment_new'), {'course':self.course.id,'due_date':'2030-01-01','due_time':'22:00','passing_score':90,'attempts_allowed':2,'departments':[self.employee.profile.department_id]})
        self.assertEqual(response.status_code, 302); self.assignment.refresh_from_db()
        self.assertEqual(self.assignment.status, 'REASSIGNED')
        latest = Assignment.objects.filter(user=self.employee, course=self.course).first()
        self.assertEqual(Assignment.objects.filter(user=self.employee, course=self.course).count(), 2)
        self.assertEqual(latest.passing_score, 90)
        self.assertEqual(latest.due_time.strftime('%H:%M'), '22:00')
        self.assertEqual(latest.attempts_allowed, 2)
    def test_successful_test_blocks_further_attempts(self):
        self.client.login(username='employee', password='test-password')
        self.client.post(reverse('test_start', args=[self.course.id]))
        response = self.client.post(reverse('test', args=[self.course.id]), {f'q{self.test.questions.first().id}':'0'})
        self.assertContains(response, 'Тест пройден'); self.assignment.refresh_from_db(); self.assertEqual(self.assignment.status, 'COMPLETED'); self.assertRedirects(self.client.get(reverse('test', args=[self.course.id])), reverse('course', args=[self.course.id]))

    def test_test_requires_a_separate_start_screen(self):
        self.test.time_limit_minutes = 10
        self.test.save(update_fields=['time_limit_minutes'])
        self.client.login(username='employee', password='test-password')
        start_page = self.client.get(reverse('test', args=[self.course.id]))
        self.assertContains(start_page, 'Начать тест')
        self.assertContains(start_page, '10 мин.')
        response = self.client.post(reverse('test_start', args=[self.course.id]))
        self.assertRedirects(response, reverse('test', args=[self.course.id]))
        questions = self.client.get(reverse('test', args=[self.course.id]))
        self.assertContains(questions, self.test.questions.first().text)

    def test_time_limit_blocks_an_expired_test_attempt(self):
        self.test.time_limit_minutes = 1
        self.test.save(update_fields=['time_limit_minutes'])
        self.client.login(username='employee', password='test-password')
        self.assignment.current_attempt_started_at = timezone.now() - timedelta(minutes=2)
        self.assignment.save(update_fields=['current_attempt_started_at'])
        response = self.client.get(reverse('test', args=[self.course.id]))
        self.assertContains(response, 'Время истекло')
        self.assertEqual(self.assignment.attempts.count(), 1)

    def test_test_timer_cannot_be_reset_or_bypassed(self):
        self.test.time_limit_minutes = 10
        self.test.save(update_fields=['time_limit_minutes'])
        self.client.login(username='employee', password='test-password')
        answer = {f'q{self.test.questions.first().id}': '0'}
        self.client.post(reverse('test', args=[self.course.id]), answer)
        self.assertFalse(self.assignment.attempts.exists())
        self.client.post(reverse('test_start', args=[self.course.id]))
        self.assignment.refresh_from_db()
        started_at = self.assignment.current_attempt_started_at
        self.assertIsNotNone(started_at)
        self.client.post(reverse('test_start', args=[self.course.id]))
        self.assignment.refresh_from_db()
        self.assertEqual(self.assignment.current_attempt_started_at, started_at)
        self.client.post(reverse('test', args=[self.course.id]), answer)
        self.client.post(reverse('test', args=[self.course.id]), answer)
        self.assertEqual(self.assignment.attempts.count(), 1)

    def test_training_history_protects_users_and_courses_from_deletion(self):
        with self.assertRaises(ProtectedError):
            self.employee.delete()
        with self.assertRaises(ProtectedError):
            self.course.delete()
        with self.assertRaises(ProtectedError):
            self.course.direction.delete()
    def test_reassignment_preserves_completed_history(self):
        self.assignment.status = 'COMPLETED'; self.assignment.completed_at = timezone.now(); self.assignment.save()
        self.client.login(username='methodist', password='test-password')
        self.client.post(reverse('assignment_new'), {'course':self.course.id,'due_date':'2030-01-01','passing_score':85,'attempts_allowed':2,'users':[self.employee.id]})
        cycles = Assignment.objects.filter(user=self.employee, course=self.course)
        self.assertEqual(cycles.count(), 2)
        self.assertEqual(cycles.filter(status='COMPLETED').count(), 1)

    def test_report_places_attempts_in_separate_columns(self):
        self.assignment.attempts_allowed = 2
        self.assignment.save(update_fields=['attempts_allowed'])
        TestAttempt.objects.create(
            assignment=self.assignment, number=1, correct_answers=8,
            score=80, passed=True,
        )
        TestAttempt.objects.create(
            assignment=self.assignment, number=2, correct_answers=10,
            score=100, passed=True,
        )
        Assignment.objects.create(
            user=self.employee, course=self.course, test_version=self.test,
            due_date=timezone.localdate() + timedelta(days=30),
            passing_score=80, attempts_allowed=2,
        )
        self.client.login(username='methodist', password='test-password')

        page = self.client.get(reverse('reports'))
        self.assertContains(page, 'Назначение 2')
        self.assertContains(page, '100%')
        self.assertNotContains(page, 'Попытка 1')

        response = self.client.get(reverse('report_excel'))
        workbook = load_workbook(BytesIO(response.content))
        sheet = workbook.active
        headers = [cell.value for cell in sheet[1]]
        first_assignment = 'Назначение 1'
        self.assertIn(first_assignment, headers)
        self.assertIn('Назначение 2', headers)
        self.assertNotIn('Попытка 1', headers)
        self.assertEqual(sheet.max_row, 2)
        values = [cell.value for cell in sheet[2]]
        self.assertIn('100%', values[headers.index(first_assignment)])

    def test_ten_reassignments_create_ten_distinct_history_cycles(self):
        self.assignment.delete()
        self.client.login(username='methodist', password='test-password')
        for day in range(1, 11):
            self.client.post(reverse('assignment_new'), {
                'course': self.course.id, 'due_date': f'2030-01-{day:02d}',
                'due_time': '22:00', 'passing_score': 80,
                'attempts_allowed': 2, 'users': [self.employee.id],
            })
        cycles = Assignment.objects.filter(user=self.employee, course=self.course)
        self.assertEqual(cycles.count(), 10)
        self.assertEqual(cycles.filter(status='REASSIGNED').count(), 9)
        self.assertEqual(cycles.filter(status='ASSIGNED').count(), 1)
    def test_multiple_answer_question_is_scored(self):
        question = Question.objects.create(test=self.test, text='Выберите два', question_type='MULTIPLE', options=['A','B','C'], correct_indexes=[0,2], points=2, order=2)
        self.assignment.attempts_allowed = 2; self.assignment.save()
        self.client.login(username='employee', password='test-password')
        self.client.post(reverse('test_start', args=[self.course.id]))
        response = self.client.post(reverse('test', args=[self.course.id]), {f'q{self.test.questions.first().id}':'0', f'q{question.id}':['0','2']})
        self.assertContains(response, '100%')
    def test_edit_after_attempt_creates_new_test_version(self):
        TestAttempt.objects.create(assignment=self.assignment, number=1, correct_answers=0, score=0, passed=False)
        self.client.login(username='methodist', password='test-password')
        response = self.client.post(reverse('test_editor', args=[self.course.id]), {'action':'save_test','title':'Новая версия','instructions':''})
        self.assertEqual(response.status_code, 302)
        self.test.refresh_from_db(); self.assertFalse(self.test.is_current)
        current = Test.objects.get(course=self.course, is_current=True)
        self.assertEqual(current.version, 2); self.assertEqual(current.status, 'DRAFT')
    def test_document_bank_available_to_employee(self):
        self.client.login(username='employee', password='test-password')
        self.assertEqual(self.client.get(reverse('document_bank')).status_code, 200)

    def test_bank_document_can_be_linked_without_duplicate_legacy_lesson(self):
        legacy = Lesson.objects.create(course=self.course, title='Старая копия', category='INTERNAL', kind='DOC')
        document = InternalDocument.objects.create(
            title='Единый документ',
            file=SimpleUploadedFile('document.pdf', b'%PDF-1.4 test', content_type='application/pdf'),
        )
        self.client.login(username='methodist', password='test-password')
        response = self.client.post(reverse('course_document_link', args=[self.course.id]), {'documents': [document.id]})
        self.assertRedirects(response, reverse('course_content', args=[self.course.id]))
        self.assertTrue(document.courses.filter(pk=self.course.pk).exists())
        response = self.client.get(reverse('course', args=[self.course.id]))
        self.assertContains(response, 'Единый документ')
        self.assertNotContains(response, legacy.title)

    def test_employee_cannot_change_course_document_links(self):
        document = InternalDocument.objects.create(
            title='Закрытый документ',
            file=SimpleUploadedFile('private.pdf', b'%PDF-1.4 test', content_type='application/pdf'),
        )
        self.client.login(username='employee', password='test-password')
        response = self.client.post(reverse('course_document_link', args=[self.course.id]), {'documents': [document.id]})
        self.assertEqual(response.status_code, 403)
        self.assertFalse(document.courses.filter(pk=self.course.pk).exists())

    def test_overdue_assignment_blocks_acknowledgement_and_test(self):
        self.assignment.due_date = timezone.localdate() - timedelta(days=1)
        self.assignment.save(update_fields=['due_date'])
        self.client.login(username='employee', password='test-password')
        response = self.client.post(reverse('acknowledge', args=[self.assignment.id]), follow=True)
        self.assertContains(response, 'Срок прохождения курса истёк')
        response = self.client.post(reverse('test', args=[self.course.id]), {f'q{self.test.questions.first().id}': '0'}, follow=True)
        self.assertContains(response, 'Срок теста истёк')
        self.assignment.refresh_from_db()
        self.assertEqual(self.assignment.status, 'OVERDUE')
        self.assertIsNone(self.assignment.acknowledged_at)
        self.assertEqual(self.assignment.attempts.count(), 0)

    def test_methodist_can_create_and_edit_course(self):
        self.client.login(username='methodist', password='test-password')
        response = self.client.post(reverse('course_new'), {
            'category': 'Новая категория',
            'title': 'Новый учебный курс',
            'description': 'Краткое описание',
            'active': 'on',
        })
        created = Course.objects.get(title='Новый учебный курс')
        self.assertEqual(created.direction.title, 'Новая категория')
        self.assertRedirects(response, reverse('course_content', args=[created.id]))
        response = self.client.post(reverse('course_edit', args=[created.id]), {
            'category': 'Обновлённая категория',
            'title': 'Обновлённый курс',
            'description': 'Новое описание',
        })
        self.assertRedirects(response, reverse('course_content', args=[created.id]))
        created.refresh_from_db()
        self.assertEqual(created.title, 'Обновлённый курс')
        self.assertFalse(created.active)

    def test_pdf_can_open_inline_or_download(self):
        document = InternalDocument.objects.create(
            title='Документ PDF',
            file=SimpleUploadedFile('manual.pdf', b'%PDF-1.4 test', content_type='application/pdf'),
        )
        document.courses.add(self.course)
        self.client.login(username='employee', password='test-password')
        page = self.client.get(reverse('course', args=[self.course.id]))
        self.assertContains(page, 'target="_blank"')
        self.assertContains(page, '?download=1')
        inline = self.client.get(document.file.url)
        download = self.client.get(document.file.url + '?download=1')
        self.assertIn('inline', inline.headers['Content-Disposition'])
        self.assertIn('attachment', download.headers['Content-Disposition'])

    def test_assigned_courses_show_only_latest_assignments_with_a_test(self):
        unassigned = Course.objects.create(
            direction=self.course.direction, title='Доступный, но не назначенный курс',
        )
        no_test_course = Course.objects.create(
            direction=self.course.direction, title='Назначение без теста',
        )
        Assignment.objects.create(
            user=self.employee, course=no_test_course,
            due_date=timezone.localdate() + timedelta(days=2),
        )
        self.client.login(username='employee', password='test-password')

        response = self.client.get(reverse('courses'))

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, self.course.title)
        self.assertContains(response, self.test.title)
        self.assertNotContains(response, unassigned.title)
        self.assertNotContains(response, no_test_course.title)
        self.assertContains(response, 'К прохождению')
        self.assertContains(response, 'Пройденные')
        self.assertEqual(response.context['active_count'], 1)
        self.assertEqual(response.context['completed_count'], 0)

    def test_completed_courses_are_shown_only_on_completed_tab(self):
        self.assignment.status = 'COMPLETED'
        self.assignment.completed_at = timezone.now()
        self.assignment.save(update_fields=['status', 'completed_at'])
        TestAttempt.objects.create(
            assignment=self.assignment, number=1, correct_answers=1, score=100, passed=True,
        )
        self.client.login(username='employee', password='test-password')

        active = self.client.get(reverse('courses'))
        completed = self.client.get(reverse('courses') + '?tab=completed')

        self.assertNotContains(active, self.course.title)
        self.assertContains(active, 'Курсов к прохождению пока нет')
        self.assertContains(completed, self.course.title)
        self.assertContains(completed, 'Последний результат: <strong>100%</strong>', html=True)
        self.assertEqual(completed.context['active_count'], 0)
        self.assertEqual(completed.context['completed_count'], 1)

    def test_reassigned_course_keeps_previous_completion_on_completed_tab(self):
        self.assignment.status = 'COMPLETED'
        self.assignment.completed_at = timezone.now() - timedelta(days=2)
        self.assignment.save(update_fields=['status', 'completed_at'])
        TestAttempt.objects.create(
            assignment=self.assignment, number=1, correct_answers=1, score=100, passed=True,
        )
        Assignment.objects.create(
            user=self.employee, course=self.course, test_version=self.test,
            due_date=timezone.localdate() + timedelta(days=3), status='ASSIGNED',
        )
        self.client.login(username='employee', password='test-password')

        active = self.client.get(reverse('courses'))
        completed = self.client.get(reverse('courses') + '?tab=completed')

        self.assertContains(active, self.course.title)
        self.assertContains(completed, self.course.title)
        self.assertEqual(active.context['active_count'], 1)
        self.assertEqual(completed.context['completed_count'], 1)

    def test_cancelled_latest_assignment_hides_course_from_assigned_list(self):
        self.assignment.status = 'REASSIGNED'
        self.assignment.save(update_fields=['status'])
        Assignment.objects.create(
            user=self.employee, course=self.course, test_version=self.test,
            due_date=timezone.localdate() + timedelta(days=3), status='CANCELLED',
        )
        self.client.login(username='employee', password='test-password')

        response = self.client.get(reverse('courses'))

        self.assertNotContains(response, self.course.title)
        self.assertContains(response, 'Курсов к прохождению пока нет')

    def test_google_drive_material_is_embedded_and_opens_in_new_tab(self):
        lesson = Lesson.objects.create(
            course=self.course, title='Материал Drive', kind='VIDEO', order=1,
            url='https://drive.google.com/file/d/test-file_123/view',
        )
        self.client.login(username='employee', password='test-password')
        response = self.client.get(reverse('course', args=[self.course.id]))
        self.assertContains(response, lesson.drive_embed_url)
        self.assertContains(response, 'Открыть в Google Drive')
        self.assertContains(response, 'target="_blank"')

    def test_test_options_keep_original_values_when_display_order_is_shuffled(self):
        self.client.login(username='employee', password='test-password')
        self.client.post(reverse('test_start', args=[self.course.id]))
        response = self.client.get(reverse('test', args=[self.course.id]))
        self.assertContains(response, 'Порядок вариантов ответа меняется')
        self.assertContains(response, 'value="0"')
        self.assertContains(response, 'value="1"')

    def test_login_shows_welcome_identity_and_thirty_minute_session(self):
        self.employee.first_name = 'Иван'
        self.employee.last_name = 'Иванов'
        self.employee.save(update_fields=['first_name', 'last_name'])
        response = self.client.post(reverse('login'), {
            'username': 'employee', 'password': 'test-password',
        }, follow=True)
        self.assertContains(response, 'Добро пожаловать, Иван Иванов!')
        self.assertContains(response, 'Логин: employee')
        self.assertEqual(settings.SESSION_COOKIE_AGE, 1800)
        self.assertContains(response, 'const limit=1800000')
        self.assertFalse(settings.SESSION_SAVE_EVERY_REQUEST)
        self.assertEqual(self.client.get(reverse('session_ping')).status_code, 204)
        self.assertGreaterEqual(self.client.session.get_expiry_age(), 1799)

    def test_load_mcbp2_is_idempotent(self):
        call_command('load_mcbp2')
        call_command('load_mcbp2')
        course = Course.objects.get(title='МЦБП 2 — Эффективная коммуникация и перевод пациентов')
        self.assertTrue(course.active)
        self.assertEqual(course.lessons.count(), 3)
        self.assertEqual(
            list(course.lessons.values_list('order', flat=True)),
            [1, 2, 3],
        )
        self.assertTrue(all(lesson.drive_embed_url for lesson in course.lessons.all()))
        test = course.tests.get(is_current=True)
        self.assertEqual(test.status, 'PUBLISHED')
        self.assertEqual(test.time_limit_minutes, 10)
        self.assertEqual(test.questions.count(), 10)
        self.assertEqual(
            list(test.questions.values_list('correct_index', flat=True)),
            [1, 1, 1, 0, 2, 1, 1, 2, 1, 1],
        )

    def test_load_mcbp346_is_idempotent_and_preserves_history(self):
        call_command('load_mcbp346', with_history=True)
        call_command('load_mcbp346', with_history=True)
        courses = Course.objects.filter(title__startswith='МЦБП ').exclude(
            title='Тестовый курс'
        ).order_by('title')
        self.assertEqual(courses.count(), 3)
        self.assertEqual(
            [course.lessons.count() for course in courses],
            [4, 5, 3],
        )
        self.assertEqual(
            [course.tests.get(is_current=True).questions.count() for course in courses],
            [10, 10, 10],
        )
        # employee and methodist are staff; sysadmin is deliberately excluded.
        history = Assignment.objects.filter(demo_key__contains='-history-202603:')
        self.assertEqual(history.count(), 6)
        self.assertEqual(TestAttempt.objects.filter(assignment__in=history).count(), 6)
        self.assertEqual(set(history.values_list('attempts__score', flat=True)), {80, 90, 100})
        self.assertFalse(history.filter(user=self.admin).exists())
        self.assertEqual(
            set(history.values_list('completed_at__date', flat=True)),
            {
                date(2026, 3, 2),
                date(2026, 3, 3),
                date(2026, 3, 4),
            },
        )

    def test_safety_and_communication_courses_are_idempotent_and_unassigned(self):
        call_command('load_safety_communication_courses')
        call_command('load_safety_communication_courses')
        communication = Course.objects.get(title='Эффективная коммуникация с пациентом')
        blue_code = Course.objects.get(title='Синий код: базовая реанимация и действия команды')
        self.assertEqual(communication.direction.title, 'Коммуникация с пациентами и профессиональное развитие')
        self.assertEqual(blue_code.direction.title, 'Коды безопасности и экстренное реагирование')
        self.assertEqual(
            communication.lessons.order_by('order', 'id').first().title,
            'Хирургическая точность в диалоге: Мастер-класс онкохирурга Вадима Гущина (Mercy Medical Center, США)',
        )
        for course in (communication, blue_code):
            self.assertTrue(course.active)
            self.assertEqual(course.lessons.count(), 3)
            self.assertTrue(all(lesson.drive_embed_url for lesson in course.lessons.all()))
            test = course.tests.get(is_current=True)
            self.assertEqual(test.status, 'PUBLISHED')
            self.assertEqual(test.time_limit_minutes, 10)
            self.assertEqual(test.questions.count(), 10)
            self.assertFalse(Assignment.objects.filter(course=course).exists())

        blue_test = blue_code.tests.get(is_current=True)
        self.assertEqual(blue_test.title, 'Итоговый тест: Синий код — базовая реанимация взрослых')
        self.assertEqual(
            list(blue_test.questions.order_by('order').values_list('text', flat=True)),
            [item[0] for item in BLUE_CODE_ADULT_QUESTIONS],
        )
        combined_text = ' '.join(blue_test.questions.values_list('text', flat=True)).casefold()
        self.assertNotIn('младен', combined_text)
        self.assertNotIn('у детей', combined_text)
        self.assertNotIn('ребен', combined_text)

    def test_replace_blue_code_test_preserves_history_and_updates_only_incomplete_cycles(self):
        call_command('load_safety_communication_courses')
        course = Course.objects.get(title='Синий код: базовая реанимация и действия команды')
        old_test = course.tests.get(is_current=True)
        old_test.title = f'Итоговый тест: {course.title}'
        old_test.save(update_fields=['title'])

        incomplete = Assignment.objects.create(
            user=self.employee, course=course, test_version=old_test,
            due_date=date(2030, 10, 15), due_time='23:59',
            passing_score=80, attempts_allowed=2, status='FAILED',
        )
        TestAttempt.objects.create(
            assignment=incomplete, number=1, correct_answers=7, score=70, passed=False,
        )
        completed = Assignment.objects.create(
            user=self.admin, course=course, test_version=old_test,
            due_date=date(2030, 10, 15), passing_score=80, attempts_allowed=1,
            status='COMPLETED', completed_at=timezone.now(),
        )
        TestAttempt.objects.create(
            assignment=completed, number=1, correct_answers=9, score=90, passed=True,
        )

        call_command('replace_blue_code_adult_test')
        call_command('replace_blue_code_adult_test')

        old_test.refresh_from_db()
        adult_test = course.tests.get(is_current=True)
        incomplete.refresh_from_db()
        completed.refresh_from_db()
        self.assertFalse(old_test.is_current)
        self.assertEqual(adult_test.version, old_test.version + 1)
        self.assertEqual(adult_test.time_limit_minutes, 10)
        self.assertEqual(adult_test.questions.count(), 10)
        self.assertEqual(
            list(adult_test.questions.order_by('order').values_list('correct_index', flat=True)),
            [item[2] for item in BLUE_CODE_ADULT_QUESTIONS],
        )
        self.assertEqual(incomplete.status, 'REASSIGNED')
        self.assertEqual(incomplete.attempts.count(), 1)
        replacement = Assignment.objects.get(
            user=self.employee, course=course, test_version=adult_test,
        )
        self.assertEqual(replacement.status, 'ASSIGNED')
        self.assertEqual(replacement.due_date, date(2030, 10, 15))
        self.assertEqual(replacement.due_time.strftime('%H:%M'), '23:59')
        self.assertEqual(replacement.passing_score, 80)
        self.assertEqual(replacement.attempts_allowed, 2)
        self.assertEqual(replacement.attempts.count(), 0)
        self.assertEqual(completed.status, 'COMPLETED')
        self.assertEqual(completed.test_version, old_test)
        self.assertEqual(completed.attempts.count(), 1)
        self.assertEqual(course.tests.count(), 2)

    def test_replace_mcbp6_test_creates_ten_question_version_for_only_incomplete_cycles(self):
        call_command('load_mcbp346')
        course = Course.objects.get(title='МЦБП 6 — Профилактика падений')
        old_test = course.tests.get(is_current=True)
        incomplete = Assignment.objects.create(
            user=self.employee, course=course, test_version=old_test,
            due_date=date(2030, 1, 1),
            passing_score=80, attempts_allowed=2, status='FAILED',
        )
        TestAttempt.objects.create(
            assignment=incomplete, number=1, correct_answers=7, score=78, passed=False,
        )
        completed = Assignment.objects.create(
            user=self.admin, course=course, test_version=old_test,
            due_date=date(2030, 1, 1), passing_score=80, attempts_allowed=1,
            status='COMPLETED', completed_at=timezone.now(),
        )
        TestAttempt.objects.create(
            assignment=completed, number=1, correct_answers=9, score=100, passed=True,
        )

        call_command('replace_mcbp6_test')
        call_command('replace_mcbp6_test')

        old_test.refresh_from_db()
        corrected = course.tests.get(is_current=True)
        incomplete.refresh_from_db()
        self.assertFalse(old_test.is_current)
        self.assertEqual(corrected.version, 2)
        self.assertEqual(corrected.questions.count(), 10)
        self.assertEqual(
            list(corrected.questions.values_list('correct_index', flat=True)),
            [1, 2, 1, 3, 0, 2, 0, 3, 1, 2],
        )
        self.assertEqual(incomplete.status, 'REASSIGNED')
        self.assertEqual(incomplete.attempts.count(), 1)
        replacement = Assignment.objects.get(user=self.employee, course=course, test_version=corrected)
        self.assertEqual(replacement.status, 'ASSIGNED')
        self.assertEqual(replacement.attempts.count(), 0)
        self.assertEqual(replacement.due_date, date(2030, 1, 1))
        completed.refresh_from_db()
        self.assertEqual(completed.status, 'COMPLETED')
        self.assertEqual(completed.test_version, old_test)

    def test_category_and_department_assignment_uses_intersection(self):
        doctor = EmployeeType.objects.get(name='Врач')
        self.employee.profile.employee_type = doctor
        self.employee.profile.save(update_fields=['employee_type'])
        nurse_type = EmployeeType.objects.get(name='Медсестра')
        other_department = Department.objects.create(name='Другое отделение')
        nurse = User.objects.create_user('nurse', password='test-password')
        Profile.objects.create(
            user=nurse, role='EMPLOYEE', department=self.employee.profile.department,
            employee_type=nurse_type, position='Медсестра',
        )
        other_doctor = User.objects.create_user('other-doctor', password='test-password')
        Profile.objects.create(
            user=other_doctor, role='EMPLOYEE', department=other_department,
            employee_type=doctor, position='Врач',
        )
        self.assignment.delete()
        self.client.login(username='methodist', password='test-password')
        response = self.client.post(reverse('assignment_new'), {
            'course': self.course.id, 'due_date': '2030-01-01',
            'due_time': '22:00', 'passing_score': 80, 'attempts_allowed': 2,
            'departments': [self.employee.profile.department_id],
            'employee_types': [doctor.id],
        })
        self.assertEqual(response.status_code, 302)
        self.assertTrue(Assignment.objects.filter(user=self.employee, course=self.course).exists())
        self.assertFalse(Assignment.objects.filter(user=nurse, course=self.course).exists())
        self.assertFalse(Assignment.objects.filter(user=other_doctor, course=self.course).exists())

    def test_category_only_assignment_targets_that_category(self):
        nurse_type = EmployeeType.objects.get(name='Медсестра')
        nurse = User.objects.create_user('nurse', password='test-password')
        Profile.objects.create(
            user=nurse, role='EMPLOYEE', department=self.employee.profile.department,
            employee_type=nurse_type, position='Медсестра',
        )
        self.assignment.delete()
        self.client.login(username='methodist', password='test-password')
        self.client.post(reverse('assignment_new'), {
            'course': self.course.id, 'due_date': '2030-01-01',
            'due_time': '22:00', 'passing_score': 80, 'attempts_allowed': 2,
            'employee_types': [nurse_type.id],
        })
        self.assertTrue(Assignment.objects.filter(user=nurse, course=self.course).exists())
        self.assertFalse(Assignment.objects.filter(user=self.employee, course=self.course).exists())

    def test_assignment_employee_search_matches_active_employee_list(self):
        inactive = User.objects.create_user('inactive-user', password='test-password', is_active=False)
        Profile.objects.create(
            user=inactive, role='EMPLOYEE', department=self.employee.profile.department,
            employee_type=self.employee.profile.employee_type, position='Врач',
        )
        self.client.login(username='methodist', password='test-password')
        response = self.client.get(reverse('assignments'))
        expected_ids = set(Profile.objects.filter(user__is_active=True).values_list('user_id', flat=True))
        displayed_ids = set(response.context['employees'].values_list('user_id', flat=True))
        self.assertEqual(displayed_ids, expected_ids)
        self.assertContains(response, 'id="employee-search"')
        self.assertContains(response, 'Выбрать найденных')
        self.assertContains(response, "option.style.display = matches ? 'block' : 'none'")
        self.assertContains(response, 'Подходящие сотрудники не найдены')
        self.assertContains(response, '(employee)')
        self.assertNotContains(response, 'inactive-user')

    def test_assignment_rejects_tampered_inactive_employee_selection(self):
        inactive = User.objects.create_user('inactive-target', password='test-password', is_active=False)
        Profile.objects.create(
            user=inactive, role='EMPLOYEE', department=self.employee.profile.department,
            employee_type=self.employee.profile.employee_type, position='Врач',
        )
        self.client.login(username='methodist', password='test-password')
        response = self.client.post(reverse('assignment_new'), {
            'course': self.course.id, 'due_date': '2030-01-01', 'due_time': '22:00',
            'passing_score': 80, 'attempts_allowed': 2, 'users': [inactive.id],
        })
        self.assertEqual(response.status_code, 400)
        self.assertFalse(Assignment.objects.filter(user=inactive, course=self.course).exists())

    def test_forgot_password_is_generic_and_does_not_duplicate_requests(self):
        existing = self.client.post(reverse('forgot_password'), {
            'username': 'employee', 'contact_method': 'WHATSAPP',
            'contact_value': '+7 777 123 45 67',
        })
        repeated = self.client.post(reverse('forgot_password'), {
            'username': 'employee', 'contact_method': 'EMAIL',
            'contact_value': 'employee@example.kz',
        })
        unknown = self.client.post(reverse('forgot_password'), {
            'username': 'does-not-exist', 'contact_method': 'EMAIL',
            'contact_value': 'unknown@example.kz',
        })
        confirmation = 'Запрос принят. Если такой активный логин существует'
        self.assertContains(existing, confirmation)
        self.assertContains(repeated, confirmation)
        self.assertContains(unknown, confirmation)
        self.assertEqual(
            PasswordResetRequest.objects.filter(user=self.employee, status='PENDING').count(),
            1,
        )
        password_request = PasswordResetRequest.objects.get(user=self.employee, status='PENDING')
        self.assertEqual(password_request.contact_method, 'EMAIL')
        self.assertEqual(password_request.contact_value, 'employee@example.kz')
        self.client.login(username='sysadmin', password='admin-password')
        self.assertContains(self.client.get(reverse('dashboard')), 'Запросы пароля: 1')
        self.assertContains(self.client.get(reverse('password_requests')), 'employee@example.kz')

    def test_forgot_password_validates_contact_without_disclosing_login(self):
        response = self.client.post(reverse('forgot_password'), {
            'username': 'employee', 'contact_method': 'WHATSAPP', 'contact_value': '123',
        })
        self.assertContains(response, 'Укажите номер WhatsApp полностью')
        self.assertFalse(PasswordResetRequest.objects.exists())

    def test_admin_can_dismiss_password_request_without_reset(self):
        password_request = PasswordResetRequest.objects.create(user=self.employee)
        self.client.login(username='sysadmin', password='admin-password')
        response = self.client.post(
            reverse('password_request_dismiss', args=[password_request.id]), follow=True,
        )
        self.assertContains(response, 'Запрос закрыт без смены пароля')
        password_request.refresh_from_db()
        self.assertEqual(password_request.status, 'DISMISSED')
        self.assertEqual(password_request.resolved_by, self.admin)
        self.assertTrue(self.employee.check_password('test-password'))

    def test_admin_resets_password_and_closes_request(self):
        password_request = PasswordResetRequest.objects.create(user=self.employee)
        self.client.login(username='sysadmin', password='admin-password')
        response = self.client.post(reverse('employee_reset_password', args=[self.employee.profile.id]), {
            'new_password1': 'temporary-2026',
            'new_password2': 'temporary-2026',
        }, follow=True)
        self.assertContains(response, 'Временный пароль установлен')
        self.employee.refresh_from_db()
        self.employee.profile.refresh_from_db()
        password_request.refresh_from_db()
        self.assertFalse(self.employee.check_password('test-password'))
        self.assertTrue(self.employee.check_password('temporary-2026'))
        self.assertTrue(self.employee.profile.force_password_change)
        self.assertEqual(password_request.status, 'COMPLETED')
        self.client.logout()
        login_response = self.client.post(reverse('login'), {
            'username': 'employee', 'password': 'temporary-2026',
        })
        self.assertRedirects(login_response, reverse('password'))

    def test_methodist_cannot_manage_employee_profiles_or_passwords(self):
        self.client.login(username='methodist', password='test-password')
        self.assertEqual(self.client.get(reverse('employees')).status_code, 403)
        self.assertEqual(
            self.client.post(reverse('employee_reset_password', args=[self.employee.profile.id]), {
                'new_password1': 'temporary-2026', 'new_password2': 'temporary-2026',
            }).status_code,
            403,
        )

    def test_active_assignment_can_be_edited_and_overdue_status_reopens(self):
        self.assignment.status = 'OVERDUE'
        self.assignment.due_date = timezone.localdate() - timedelta(days=1)
        self.assignment.save(update_fields=['status', 'due_date'])
        self.client.login(username='methodist', password='test-password')
        edit_page = self.client.get(reverse('assignment_edit', args=[self.assignment.id]))
        self.assertEqual(edit_page.status_code, 200)
        self.assertContains(edit_page, 'Редактирование назначения')
        new_due_date = timezone.localdate() + timedelta(days=30)
        response = self.client.post(reverse('assignment_edit', args=[self.assignment.id]), {
            'due_date': new_due_date.isoformat(), 'due_time': '21:30',
            'passing_score': 90, 'attempts_allowed': 3,
        }, follow=True)
        self.assertContains(response, 'Назначение обновлено')
        self.assignment.refresh_from_db()
        self.assertEqual(self.assignment.due_date, new_due_date)
        self.assertEqual(self.assignment.due_time.strftime('%H:%M'), '21:30')
        self.assertEqual(self.assignment.passing_score, 90)
        self.assertEqual(self.assignment.attempts_allowed, 3)
        self.assertEqual(self.assignment.status, 'ASSIGNED')

    def test_started_assignment_keeps_passing_score_when_edited(self):
        TestAttempt.objects.create(
            assignment=self.assignment, number=1, correct_answers=0,
            score=40, passed=False,
        )
        self.assignment.status = 'IN_PROGRESS'
        self.assignment.attempts_allowed = 3
        self.assignment.save(update_fields=['status', 'attempts_allowed'])
        self.client.login(username='methodist', password='test-password')
        response = self.client.post(reverse('assignment_edit', args=[self.assignment.id]), {
            'due_date': (timezone.localdate() + timedelta(days=10)).isoformat(),
            'due_time': '23:00', 'passing_score': 90, 'attempts_allowed': 3,
        }, follow=True)
        self.assertContains(response, 'Порог нельзя изменить после начала тестирования')
        self.assignment.refresh_from_db()
        self.assertEqual(self.assignment.passing_score, 80)

    def test_cancelling_assignment_preserves_history_and_blocks_test(self):
        self.client.login(username='methodist', password='test-password')
        response = self.client.post(
            reverse('assignment_cancel', args=[self.assignment.id]), follow=True,
        )
        self.assertContains(response, 'Назначение отменено. Запись сохранена в журнале назначений')
        self.assignment.refresh_from_db()
        self.assertEqual(self.assignment.status, 'CANCELLED')
        self.assertIsNotNone(self.assignment.cancelled_at)
        self.assertEqual(self.assignment.cancelled_by.username, 'methodist')
        self.client.logout()
        self.client.login(username='employee', password='test-password')
        self.client.post(reverse('acknowledge', args=[self.assignment.id]))
        self.assignment.refresh_from_db()
        self.assertEqual(self.assignment.status, 'CANCELLED')
        course_page = self.client.get(reverse('course', args=[self.course.id]))
        self.assertContains(course_page, 'Тест открывается после назначения курса')
        self.assertEqual(self.client.get(reverse('test', args=[self.course.id])).status_code, 403)

    def test_completed_assignment_cannot_be_edited_or_cancelled(self):
        self.assignment.status = 'COMPLETED'
        self.assignment.completed_at = timezone.now()
        self.assignment.save(update_fields=['status', 'completed_at'])
        self.client.login(username='methodist', password='test-password')
        edit = self.client.get(reverse('assignment_edit', args=[self.assignment.id]), follow=True)
        self.assertContains(edit, 'Можно редактировать только активное назначение')
        cancel = self.client.post(reverse('assignment_cancel', args=[self.assignment.id]), follow=True)
        self.assertContains(cancel, 'Можно отменить только активное назначение')
        self.assignment.refresh_from_db()
        self.assertEqual(self.assignment.status, 'COMPLETED')

    def test_employee_cannot_edit_or_cancel_assignment(self):
        self.client.login(username='employee', password='test-password')
        self.assertEqual(self.client.get(reverse('assignment_edit', args=[self.assignment.id])).status_code, 403)
        self.assertEqual(self.client.post(reverse('assignment_cancel', args=[self.assignment.id])).status_code, 403)

    def test_cancelled_assignments_are_excluded_from_report_and_excel(self):
        self.assignment.status = 'CANCELLED'
        self.assignment.save(update_fields=['status'])
        self.client.login(username='methodist', password='test-password')
        page = self.client.get(reverse('reports'))
        self.assertEqual(page.context['report_rows'], [])
        workbook = load_workbook(BytesIO(self.client.get(reverse('report_excel')).content))
        self.assertEqual(workbook.active.max_row, 1)

    def test_dashboard_shows_latest_course_cycle_and_employee_drilldown(self):
        self.assignment.status = 'REASSIGNED'
        self.assignment.save(update_fields=['status'])
        latest = Assignment.objects.create(
            user=self.employee, course=self.course, test_version=self.test,
            due_date=timezone.localdate() + timedelta(days=30),
            passing_score=80, attempts_allowed=2,
        )
        self.client.login(username='methodist', password='test-password')
        page = self.client.get(reverse('dashboard'), {'course': self.course.id})
        self.assertEqual(page.context['assignment_count'], 1)
        self.assertEqual(page.context['people_with_courses'], 1)
        self.assertEqual(page.context['statuses']['ASSIGNED'], 1)
        self.assertEqual(page.context['selected_row']['total'], 1)
        self.assertEqual(page.context['selected_row']['assignments'][0]['assignment'].id, latest.id)
        self.assertContains(page, self.employee.username)

    def test_dashboard_omits_cancelled_latest_cycle(self):
        self.assignment.status = 'COMPLETED'
        self.assignment.save(update_fields=['status'])
        Assignment.objects.create(
            user=self.employee, course=self.course, test_version=self.test,
            due_date=timezone.localdate() + timedelta(days=30),
            passing_score=80, attempts_allowed=2, status='CANCELLED',
        )
        self.client.login(username='methodist', password='test-password')
        page = self.client.get(reverse('dashboard'))
        self.assertEqual(page.context['assignment_count'], 0)
        self.assertEqual(page.context['statuses']['COMPLETED'], 0)

    def test_head_dashboard_only_includes_own_department(self):
        other_department = Department.objects.create(name='Иное отделение для аналитики')
        outsider = User.objects.create_user('dashboard-outsider', password='test-password')
        Profile.objects.create(
            user=outsider, role='EMPLOYEE', department=other_department,
            employee_type=self.employee.profile.employee_type, position='Врач',
        )
        Assignment.objects.create(
            user=outsider, course=self.course, test_version=self.test,
            due_date=timezone.localdate() + timedelta(days=30),
        )
        head = User.objects.create_user('dashboard-head', password='test-password')
        Profile.objects.create(
            user=head, role='HEAD', department=self.employee.profile.department,
            employee_type=self.employee.profile.employee_type, position='Заведующий',
        )
        self.client.login(username='dashboard-head', password='test-password')
        page = self.client.get(reverse('dashboard'), {'course': self.course.id})
        self.assertEqual(page.context['assignment_count'], 1)
        self.assertContains(page, self.employee.username)
        self.assertNotContains(page, outsider.username)

    def test_bulk_edit_uses_department_and_category_intersection(self):
        nurse_type = EmployeeType.objects.create(name='Медсестра для массового теста')
        nurse = User.objects.create_user('bulk-nurse', password='test-password')
        Profile.objects.create(
            user=nurse, role='EMPLOYEE', department=self.employee.profile.department,
            employee_type=nurse_type, position='Медсестра',
        )
        nurse_assignment = Assignment.objects.create(
            user=nurse, course=self.course, test_version=self.test,
            due_date=timezone.localdate() + timedelta(days=1), passing_score=80,
            attempts_allowed=1,
        )
        other_department = Department.objects.create(name='Другое массовое отделение')
        other_doctor = User.objects.create_user('bulk-other-doctor', password='test-password')
        Profile.objects.create(
            user=other_doctor, role='EMPLOYEE', department=other_department,
            employee_type=self.employee.profile.employee_type, position='Врач',
        )
        other_assignment = Assignment.objects.create(
            user=other_doctor, course=self.course, test_version=self.test,
            due_date=timezone.localdate() + timedelta(days=1), passing_score=80,
            attempts_allowed=1,
        )
        self.client.login(username='methodist', password='test-password')
        due_date = timezone.localdate() + timedelta(days=20)
        response = self.client.post(reverse('assignment_bulk'), {
            'action': 'edit', 'course': self.course.id,
            'departments': [self.employee.profile.department_id],
            'employee_types': [self.employee.profile.employee_type_id],
            'due_date': due_date.isoformat(), 'due_time': '20:30',
            'passing_score': 90, 'attempts_allowed': 2,
        }, follow=True)
        self.assertContains(response, 'Обновлено назначений: 1')
        self.assignment.refresh_from_db()
        nurse_assignment.refresh_from_db()
        other_assignment.refresh_from_db()
        self.assertEqual(self.assignment.passing_score, 90)
        self.assertEqual(self.assignment.due_time.strftime('%H:%M'), '20:30')
        self.assertEqual(nurse_assignment.passing_score, 80)
        self.assertEqual(other_assignment.passing_score, 80)

    def test_bulk_cancel_by_category_preserves_other_categories(self):
        nurse_type = EmployeeType.objects.create(name='Медсестра для отмены')
        nurse = User.objects.create_user('cancel-nurse', password='test-password')
        Profile.objects.create(
            user=nurse, role='EMPLOYEE', department=self.employee.profile.department,
            employee_type=nurse_type, position='Медсестра',
        )
        nurse_assignment = Assignment.objects.create(
            user=nurse, course=self.course, test_version=self.test,
            due_date=timezone.localdate() + timedelta(days=1), passing_score=80,
            attempts_allowed=1,
        )
        self.client.login(username='methodist', password='test-password')
        response = self.client.post(reverse('assignment_bulk'), {
            'action': 'cancel', 'course': self.course.id,
            'employee_types': [self.employee.profile.employee_type_id],
        }, follow=True)
        self.assertContains(response, 'Отменено назначений: 1')
        self.assignment.refresh_from_db()
        nurse_assignment.refresh_from_db()
        self.assertEqual(self.assignment.status, 'CANCELLED')
        self.assertEqual(nurse_assignment.status, 'ASSIGNED')

    def test_bulk_edit_is_atomic_when_an_attempt_blocks_threshold_change(self):
        colleague = User.objects.create_user('bulk-colleague', password='test-password')
        Profile.objects.create(
            user=colleague, role='EMPLOYEE', department=self.employee.profile.department,
            employee_type=self.employee.profile.employee_type, position='Врач',
        )
        colleague_assignment = Assignment.objects.create(
            user=colleague, course=self.course, test_version=self.test,
            due_date=timezone.localdate() + timedelta(days=1), passing_score=80,
            attempts_allowed=2,
        )
        TestAttempt.objects.create(
            assignment=self.assignment, number=1, correct_answers=0,
            score=40, passed=False,
        )
        self.assignment.attempts_allowed = 2
        self.assignment.save(update_fields=['attempts_allowed'])
        self.client.login(username='methodist', password='test-password')
        response = self.client.post(reverse('assignment_bulk'), {
            'action': 'edit', 'course': self.course.id,
            'employee_types': [self.employee.profile.employee_type_id],
            'due_date': (timezone.localdate() + timedelta(days=30)).isoformat(),
            'due_time': '19:00', 'passing_score': 90, 'attempts_allowed': 3,
        }, follow=True)
        self.assertContains(response, 'Массовое изменение отменено')
        self.assignment.refresh_from_db()
        colleague_assignment.refresh_from_db()
        self.assertEqual(self.assignment.passing_score, 80)
        self.assertEqual(colleague_assignment.passing_score, 80)
        self.assertEqual(colleague_assignment.attempts_allowed, 2)

    def test_bulk_action_requires_audience_and_is_forbidden_to_employee(self):
        self.client.login(username='methodist', password='test-password')
        response = self.client.post(reverse('assignment_bulk'), {
            'action': 'cancel', 'course': self.course.id,
        }, follow=True)
        self.assertContains(response, 'Выберите хотя бы одну категорию или отделение')
        self.assignment.refresh_from_db()
        self.assertEqual(self.assignment.status, 'ASSIGNED')
        self.client.logout()
        self.client.login(username='employee', password='test-password')
        self.assertEqual(self.client.post(reverse('assignment_bulk'), {
            'action': 'cancel', 'course': self.course.id,
            'employee_types': [self.employee.profile.employee_type_id],
        }).status_code, 403)

    def test_bulk_extension_supports_multiple_courses_and_reopens_failed(self):
        TestAttempt.objects.create(
            assignment=self.assignment, number=1, correct_answers=0,
            score=40, passed=False,
        )
        self.assignment.status = 'FAILED'
        self.assignment.save(update_fields=['status'])
        second_course = Course.objects.create(
            direction=self.course.direction, title='Второй тестовый курс',
        )
        second_test = Test.objects.create(
            course=second_course, title='Второй итоговый тест', status='PUBLISHED',
        )
        Question.objects.create(
            test=second_test, text='Второй вопрос?', options=['Да', 'Нет'],
            correct_index=0, correct_indexes=[0],
        )
        second_assignment = Assignment.objects.create(
            user=self.employee, course=second_course, test_version=second_test,
            due_date=timezone.localdate() - timedelta(days=1),
            status='OVERDUE', passing_score=80, attempts_allowed=2,
        )
        completed_user = User.objects.create_user('completed-user', password='test-password')
        Profile.objects.create(
            user=completed_user, role='EMPLOYEE',
            department=self.employee.profile.department,
            employee_type=self.employee.profile.employee_type, position='Врач',
        )
        completed = Assignment.objects.create(
            user=completed_user, course=self.course, test_version=self.test,
            due_date=timezone.localdate() - timedelta(days=1),
            status='COMPLETED', completed_at=timezone.now(),
            passing_score=80, attempts_allowed=1,
        )
        self.client.login(username='methodist', password='test-password')
        new_due_date = timezone.localdate() + timedelta(days=10)
        response = self.client.post(reverse('assignment_bulk'), {
            'action': 'edit', 'courses': [self.course.id, second_course.id],
            'all_audience': 'on', 'only_incomplete': 'on',
            'due_date': new_due_date.isoformat(), 'due_time': '23:59',
            'additional_attempts': 1,
        }, follow=True)
        self.assertContains(response, 'Обновлено назначений: 2')
        self.assignment.refresh_from_db()
        second_assignment.refresh_from_db()
        completed.refresh_from_db()
        self.assertEqual(self.assignment.status, 'ASSIGNED')
        self.assertEqual(self.assignment.attempts_allowed, 2)
        self.assertEqual(second_assignment.status, 'ASSIGNED')
        self.assertEqual(second_assignment.attempts_allowed, 3)
        self.assertEqual(self.assignment.due_date, new_due_date)
        self.assertEqual(second_assignment.due_date, new_due_date)
        self.assertEqual(completed.status, 'COMPLETED')
        self.assertNotEqual(completed.due_date, new_due_date)

    def test_bulk_extension_changes_only_latest_course_cycle(self):
        old_due_date = timezone.localdate() - timedelta(days=60)
        self.assignment.due_date = old_due_date
        self.assignment.status = 'FAILED'
        self.assignment.save(update_fields=['due_date', 'status'])
        latest = Assignment.objects.create(
            user=self.employee, course=self.course, test_version=self.test,
            due_date=timezone.localdate() - timedelta(days=1),
            status='OVERDUE', passing_score=80, attempts_allowed=2,
        )
        self.client.login(username='methodist', password='test-password')
        new_due_date = timezone.localdate() + timedelta(days=10)
        response = self.client.post(reverse('assignment_bulk'), {
            'action': 'edit', 'courses': [self.course.id],
            'all_audience': 'on', 'only_incomplete': 'on',
            'due_date': new_due_date.isoformat(), 'due_time': '23:59',
            'additional_attempts': 0,
        }, follow=True)
        self.assertContains(response, 'Обновлено назначений: 1')
        self.assignment.refresh_from_db()
        latest.refresh_from_db()
        self.assertEqual(self.assignment.due_date, old_due_date)
        self.assertEqual(self.assignment.status, 'FAILED')
        self.assertEqual(latest.due_date, new_due_date)
        self.assertEqual(latest.status, 'ASSIGNED')

    def test_admin_can_create_rename_and_merge_departments(self):
        self.client.login(username='sysadmin', password='admin-password')
        self.client.post(reverse('department_create'), {'name': 'Новое подразделение'})
        source = Department.objects.get(name='Новое подразделение')
        self.client.post(reverse('department_edit', args=[source.id]), {
            'name': 'Переименованное подразделение',
        })
        source.refresh_from_db()
        self.assertEqual(source.name, 'Переименованное подразделение')
        target = Department.objects.create(name='Целевое подразделение')
        colleague = User.objects.create_user('department-colleague', password='test-password')
        colleague_profile = Profile.objects.create(
            user=colleague, role='EMPLOYEE', department=source,
            employee_type=self.employee.profile.employee_type, position='Врач',
        )
        self.client.post(reverse('department_merge', args=[source.id]), {
            'target': target.id,
        })
        colleague_profile.refresh_from_db()
        self.assertEqual(colleague_profile.department, target)
        self.assertFalse(Department.objects.filter(pk=source.id).exists())

    def test_methodist_cannot_manage_departments(self):
        self.client.login(username='methodist', password='test-password')
        self.assertEqual(self.client.get(reverse('departments')).status_code, 403)
        self.assertEqual(self.client.post(reverse('department_create'), {
            'name': 'Запрещённое подразделение',
        }).status_code, 403)
