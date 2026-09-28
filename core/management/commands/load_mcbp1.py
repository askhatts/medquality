from datetime import datetime
from pathlib import Path
import random
import secrets
import string

from django.contrib.auth.models import User
from django.core.management.base import BaseCommand, CommandError
from django.db import transaction
from django.utils import timezone
from openpyxl import Workbook, load_workbook

from core.models import (
    Assignment, Course, Department, EmployeeType, InternalDocument, Lesson,
    Profile, QualityDirection, Question, Test, TestAttempt,
)


MATERIALS = [
    ('Идентификация личности пациентов', 'LINK', 'https://drive.google.com/file/d/1F0uZSE5dtKGTBTa9-FDiCwKqO3l1Mx1G/view'),
    ('Правила идентификации пациента', 'VIDEO', 'https://drive.google.com/file/d/1-HfLGPM1sWicQaLZtaGCHxjAsNRjBedb/view'),
    ('МЦБП 1 — часть 1', 'VIDEO', 'https://drive.google.com/file/d/191wi8o8MbVDiPrmFNjihc6B8Tp5wMcdR/view'),
    ('МЦБП 1 — часть 2', 'VIDEO', 'https://drive.google.com/file/d/1E1cokhiFpKlYi3EISRtPx6ghSaNrbEfW/view'),
]

QUESTIONS = [
    ('Где, согласно правилам, обязательно указываются два идентификатора пациента — Ф.И.О. и полная дата рождения?', ['На идентификационном браслете, стикере или наклейке и на каждом листе медицинской карты стационарного больного', 'Только на титульном листе медицинской карты', 'Только в электронном журнале отделения', 'Только на идентификационном браслете'], 0),
    ('Какие два основных идентификатора пациента должны использоваться на территории медицинской организации?', ['Фамилия, имя и диагноз пациента', 'Фамилия, инициалы и номер медицинской карты', 'Фамилия, имя, отчество и число, месяц, год рождения', 'Фамилия, имя, отчество и номер палаты'], 2),
    ('Медицинская сестра подошла к пациенту для введения назначенного лекарственного средства. Какое действие должно предшествовать введению препарата?', ['Сверить только номер палаты и койки', 'Уточнить только фамилию пациента', 'Провести идентификацию пациента по установленным идентификаторам', 'Проверить идентификацию после введения препарата'], 2),
    ('Как правильно провести идентификацию пациента, который может самостоятельно сообщить свои данные?', ['Попросить пациента назвать Ф.И.О. и дату рождения и сверить их с браслетом или медицинской документацией', 'Использовать только данные, сообщенные сопровождающим лицом', 'Назвать пациенту его Ф.И.О. и попросить подтвердить ответом «да»', 'Сверить только фамилию на браслете с номером палаты'], 0),
    ('Пациент находится в коматозном состоянии и не может ответить на вопросы. Как следует провести идентификацию?', ['Провести процедуру по назначению без дополнительной проверки', 'Отложить идентификацию до восстановления сознания', 'Использовать только номер палаты', 'Сверить данные идентификационного браслета с медицинской документацией'], 3),
    ('У госпитализированного пациента поврежден идентификационный браслет. Каковы правильные действия медицинской сестры?', ['Снять браслет и до выписки идентифицировать пациента только по номеру палаты', 'В кратчайшие сроки заказать дубликат, сообщить лечащему врачу, а до получения нового браслета устно уточнять Ф.И.О. и дату рождения и сверять их с медицинской документацией', 'Заменить браслет на лист бумаги с номером медицинской карты', 'Оставить поврежденный браслет, если фамилию еще можно прочитать'], 1),
    ('У пациента возникла аллергическая реакция на материал идентификационного браслета. Как следует поступить?', ['Заменить браслет на бейдж с Ф.И.О. и датой рождения и обучить пациента постоянно носить или иметь его при себе при манипуляциях', 'Нанести Ф.И.О. пациента на кожу маркером', 'Использовать только номер медицинской карты', 'Полностью отказаться от средств идентификации до выписки'], 0),
    ('Перед процедурой медицинский работник обнаружил, что манипуляцию едва не начали выполнять другому пациенту, но ошибку вовремя предотвратили. Что предусмотрено правилами?', ['Не регистрировать случай, так как вред пациенту не причинен', 'Зафиксировать случай только в медицинской карте пациента', 'Сообщить только устно заведующему отделением без оформления инцидента', 'Оформить отчет об инциденте, поскольку регистрируются также почти-ошибки'], 3),
    ('Какие данные используются для идентификации новорожденного согласно правилам?', ['Фамилия, диагноз и номер медицинской карты', 'Имя ребенка, масса тела и время рождения', 'Фамилия матери и номер палаты', 'Фамилия, число, месяц и год рождения, пол'], 3),
    ('Какой набор данных предусмотрен для идентификации неизвестного пациента?', ['Только номер медицинской карты', 'Пол, диагноз и время поступления', 'Предполагаемые Ф.И.О., возраст и номер палаты', 'Обозначение «неизвестный», признак по расе и номер медицинской карты стационарного больного'], 3),
]


class Command(BaseCommand):
    help = 'Заменяет демо-данные реальными сотрудниками и создаёт боевой курс МЦБП 1.'

    def add_arguments(self, parser):
        parser.add_argument('--doctors', required=True)
        parser.add_argument('--nurses', required=True)
        parser.add_argument('--credentials', required=True)
        parser.add_argument('--confirm-replace', action='store_true')

    @staticmethod
    def rows(path):
        workbook = load_workbook(path, read_only=True, data_only=True)
        sheet = workbook.active
        rows = sheet.iter_rows(values_only=True)
        headers = [str(value or '').strip().lower() for value in next(rows)]
        required = {'username', 'firstname', 'lastname'}
        if not required.issubset(headers):
            raise CommandError(f'{path}: обязательные столбцы {sorted(required)} не найдены.')
        for values in rows:
            row = dict(zip(headers, values))
            if row.get('username'):
                yield {key: str(row.get(key) or '').strip() for key in required}

    @staticmethod
    def password():
        alphabet = string.ascii_letters + string.digits
        return ''.join(secrets.choice(alphabet) for _ in range(12))

    def handle(self, *args, **options):
        if not options['confirm_replace']:
            raise CommandError('Добавьте --confirm-replace: операция удаляет всех пользователей, кроме sysadmin.')
        sources = [
            (Path(options['doctors']), 'Врач', 'Врач'),
            (Path(options['nurses']), 'Средний медицинский персонал', 'Медицинская сестра'),
        ]
        for path, _, _ in sources:
            if not path.is_file():
                raise CommandError(f'Файл не найден: {path}')

        imported = []
        seen = set()
        for path, employee_type, position in sources:
            for row in self.rows(path):
                username = row['username'].lower()
                if username in seen or username == 'sysadmin':
                    raise CommandError(f'Повторяющийся или зарезервированный логин: {username}')
                seen.add(username)
                imported.append((row, employee_type, position))

        credentials_path = Path(options['credentials'])
        credentials_path.parent.mkdir(parents=True, exist_ok=True)
        rng = random.Random(27022026)
        completed_at = timezone.make_aware(datetime(2026, 2, 27, 12, 0))

        with transaction.atomic():
            User.objects.exclude(username='sysadmin').delete()
            # У sysadmin также могли остаться демонстрационные назначения,
            # защищающие старые версии тестов через Assignment.test_version.
            Assignment.objects.all().delete()
            InternalDocument.objects.all().delete()
            QualityDirection.objects.all().delete()
            Department.objects.all().delete()
            EmployeeType.objects.all().delete()

            department = Department.objects.create(name='Подразделение не указано')
            types = {
                name: EmployeeType.objects.create(name=name)
                for name in ('Врач', 'Средний медицинский персонал')
            }
            direction = QualityDirection.objects.create(
                title='Безопасность пациента',
                description='Обучение безопасной идентификации пациентов перед оказанием медицинской помощи.',
                active=True,
            )
            course = Course.objects.create(
                direction=direction,
                title='МЦБП 1 — Идентификация пациента',
                description='Правила применения идентификаторов пациента и предупреждение ошибок.',
                active=True,
            )
            for order, (title, kind, url) in enumerate(MATERIALS, 1):
                Lesson.objects.create(course=course, title=title, kind=kind, category='LESSON', order=order, url=url)

            test = Test.objects.create(
                course=course, title='Итоговый тест: идентификация пациента',
                instructions='Выберите один правильный ответ. Варианты перемешиваются автоматически.',
                version=1, status='PUBLISHED', is_current=True,
            )
            for order, (text, answers, correct) in enumerate(QUESTIONS, 1):
                Question.objects.create(
                    test=test, text=text, question_type='SINGLE', options=answers,
                    correct_index=correct, correct_indexes=[correct], points=1, order=order,
                )

            output = Workbook()
            sheet = output.active
            sheet.title = 'Доступы'
            sheet.append(['Логин', 'Временный пароль', 'Фамилия', 'Имя / инициалы', 'Тип', 'Должность'])
            for row, employee_type, position in imported:
                temporary_password = self.password()
                user = User.objects.create_user(
                    username=row['username'].lower(), password=temporary_password,
                    first_name=row['firstname'], last_name=row['lastname'], email='',
                )
                Profile.objects.create(
                    user=user, role='EMPLOYEE', employee_type=types[employee_type],
                    department=department, position=position, force_password_change=True,
                )
                score = rng.randint(80, 100)
                assignment = Assignment.objects.create(
                    user=user, course=course, test_version=test,
                    due_date=completed_at.date(), passing_score=80, attempts_allowed=3,
                    status='COMPLETED', acknowledged_at=completed_at,
                    completed_at=completed_at, is_demo=True,
                    demo_key=f'mcbp1-history-{user.username}',
                )
                Assignment.objects.filter(pk=assignment.pk).update(assigned_at=completed_at)
                attempt = TestAttempt.objects.create(
                    assignment=assignment, number=1, correct_answers=round(score / 10),
                    score=score, passed=True, answers_snapshot=[],
                )
                TestAttempt.objects.filter(pk=attempt.pk).update(completed_at=completed_at)
                sheet.append([user.username, temporary_password, user.last_name, user.first_name, employee_type, position])
            output.save(credentials_path)

        self.stdout.write(self.style.SUCCESS(
            f'Импортировано {len(imported)} сотрудников; создан курс, тест и история на 27.02.2026. '
            f'Доступы: {credentials_path}'
        ))
